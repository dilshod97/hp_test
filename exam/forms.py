from datetime import date, datetime
from decimal import Decimal

from django import forms
from django.contrib.auth.forms import AuthenticationForm, SetPasswordForm

from .models import Block, Candidate, Kind, Question, QuestionBank, Test, User
from .utils import clean_html, is_empty_html


class StaffLoginForm(AuthenticationForm):
    error_messages = {
        "invalid_login": "Логин ёки парол нотўғри.",
        "inactive": "Бу фойдаланувчи блокланган.",
    }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["username"].label = "Логин"
        self.fields["password"].label = "Парол"


class CandidateLoginForm(forms.Form):
    login = forms.CharField(label="Логин", max_length=12,
                            widget=forms.TextInput(attrs={"inputmode": "numeric", "autocomplete": "off", "autofocus": True}))
    password = forms.CharField(label="Парол", max_length=12,
                               widget=forms.PasswordInput(attrs={"inputmode": "numeric", "autocomplete": "off"}))

    def clean_login(self):
        return "".join(self.cleaned_data["login"].split())

    def clean_password(self):
        return "".join(self.cleaned_data["password"].split())


def parse_birth(value):
    """«1995», «12.04.1995», «1995-04-12» ёки Excel санаси → (date ёки None, йил)."""
    if isinstance(value, (int, float)) and 1940 <= int(value) <= date.today().year:
        return None, int(value)
    d = parse_date(value)
    if d:
        return d, d.year
    text = str(value or "").strip()
    if text.isdigit() and len(text) == 4:
        return None, int(text)
    return None, None


def check_birth(born_date, year):
    if year is None:
        raise forms.ValidationError("Туғилган йилни (масалан, 1995) ёки санасини (12.04.1995) ёзинг.")
    if year < 1940 or year > date.today().year - 16:
        raise forms.ValidationError("Туғилган йил нотўғри.")
    return born_date, year


class CandidateForm(forms.Form):
    """Номзодни қўлда рўйхатга қўшиш (HR)."""
    last_name = forms.CharField(label="Фамилия", max_length=100)
    first_name = forms.CharField(label="Исм", max_length=100)
    middle_name = forms.CharField(label="Отасининг исми", max_length=100, required=False)
    birth = forms.CharField(label="Туғилган йили", max_length=10,
                            widget=forms.TextInput(attrs={"placeholder": "1995 ёки 12.04.1995"}))
    test = forms.ModelChoiceField(label="Тест", queryset=Test.objects.all(), empty_label="— тестни танланг —")

    def clean_birth(self):
        return check_birth(*parse_birth(self.cleaned_data["birth"]))


class CandidateImportForm(forms.Form):
    file = forms.FileField(
        label="Excel файл (.xlsx)",
        help_text="Устунлар: Фамилия, Исм, Отасининг исми, Туғилган йили (1995) ёки санаси (12.04.1995). Биринчи қатор — сарлавҳа.",
    )
    test = forms.ModelChoiceField(label="Тест", queryset=Test.objects.all(), empty_label="— тестни танланг —")

    def clean_file(self):
        from openpyxl import load_workbook

        f = self.cleaned_data["file"]
        if not f.name.lower().endswith(".xlsx"):
            raise forms.ValidationError("Фақат .xlsx файл қабул қилинади.")
        try:
            ws = load_workbook(f, read_only=True, data_only=True).active
        except Exception:
            raise forms.ValidationError("Файлни ўқиб бўлмади.")
        rows, errors = [], []
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cells = [c for c in (row or [])][:4] + [None] * 4
            last, first, middle, born = cells[:4]
            if not any([last, first, born]):
                continue
            born_date, year = parse_birth(born)
            if not last or not first or year is None:
                errors.append(f"{n}-қатор: фамилия, исм ёки туғилган йил нотўғри")
                continue
            rows.append((str(last).strip(), str(first).strip(), str(middle or "").strip(), (born_date, year)))
        if errors:
            raise forms.ValidationError(errors[:10])
        if not rows:
            raise forms.ValidationError("Файлда номзодлар топилмади.")
        self.rows = rows
        return f


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


class TestForm(forms.ModelForm):
    class Meta:
        model = Test
        fields = [
            "title", "description", "is_active", "free_navigation", "require_photo", "shuffle_questions", "shuffle_options", "review_mode",
            "interview_questions", "weight_practical", "weight_logic", "weight_interview", "weight_defense",
        ]
        widgets = {
            "description": forms.Textarea(attrs={"rows": 3}),
            "interview_questions": forms.CheckboxSelectMultiple,
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["interview_questions"].queryset = (
            Question.objects.filter(bank__kind=Kind.INTERVIEW).select_related("bank").order_by("bank__order", "order")
        )
        if self.instance.pk and self.instance.is_locked:
            # Номзодлар бошлаган тестда вақт тартибини ўзгартириб бўлмайди
            self.fields["free_navigation"].disabled = True
            self.fields["free_navigation"].help_text = "Номзодлар тестни бошлаган — бу созламани ўзгартириб бўлмайди."

    def clean(self):
        data = super().clean()
        total = sum(data.get(f) or 0 for f in ["weight_practical", "weight_logic", "weight_interview", "weight_defense"])
        if total != 100:
            self.add_error(None, f"Вазнлар йиғиндиси 100 бўлиши керак (ҳозир {total:g}).")
        return data


class TestCheckersForm(forms.ModelForm):
    class Meta:
        model = Test
        fields = ["checkers"]
        widgets = {"checkers": forms.CheckboxSelectMultiple}


class BlockForm(forms.ModelForm):
    class Meta:
        model = Block
        fields = ["title", "kind", "description", "duration_minutes", "order"]
        widgets = {
            "title": forms.TextInput(attrs={"placeholder": "Масалан: Мантиқий фикрлаш тести"}),
            "description": forms.Textarea(attrs={"rows": 3, "placeholder": "Номзод блок бошида шу матнни кўради"}),
        }
        help_texts = {"order": "Блоклар шу тартибда кўрсатилади (1, 2, 3...)"}

    def clean_duration_minutes(self):
        m = self.cleaned_data["duration_minutes"]
        if not 1 <= m <= 600:
            raise forms.ValidationError("1 дан 600 дақиқагача бўлиши керак.")
        return m


class BankForm(forms.ModelForm):
    class Meta:
        model = QuestionBank
        fields = ["name", "kind", "source", "description", "order"]
        labels = {"order": "Тартиб"}
        widgets = {"description": forms.Textarea(attrs={"rows": 3})}


def parse_criteria(text):
    """Ҳар қаторда битта мезон. Охирида "| 3" бўлса — макс. балл 3, акс ҳолда 2."""
    result = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        score = 2
        if "|" in line:
            body, _, tail = line.rpartition("|")
            if tail.strip().isdigit():
                line, score = body.strip(), int(tail.strip())
        result.append((line, score))
    return result


class QuestionForm(forms.ModelForm):
    use_required_attribute = False  # TinyMCE yashirin textarea bilan ishlaydi

    options_text = forms.CharField(
        label="Вариантлар (ҳар бири янги қаторда)", required=False,
        widget=forms.Textarea(attrs={"rows": 5, "placeholder": "Биринчи вариант\nИккинчи вариант\n..."}),
        help_text="Номзодга А, Б, В, Г, Д ҳарфлари билан кўрсатилади.",
    )
    correct_number = forms.IntegerField(
        label="Тўғри вариант рақами (ихтиёрий)", required=False, min_value=1,
        help_text="Берилса, савол автоматик баҳоланади (1 = А, 2 = Б, ...). Бўш қолса — эксперт баҳолайди.",
    )
    criteria_text = forms.CharField(
        label="Баҳолаш мезонлари (ҳар бири янги қаторда)", required=False,
        widget=forms.Textarea(attrs={"rows": 6, "placeholder": "Банк қолдиғини тўғри тузатади.\nРегистр қолдиғини тўғри тузатади."}),
        help_text="Ҳар бир мезон 0–2 балл (бошқа балл учун охирига \"| 3\" ёзинг). Мезонлар берилса, максимал балл уларнинг йиғиндиси бўлади. Номзодга кўрсатилмайди.",
    )

    class Meta:
        model = Question
        fields = ["title", "text", "qtype", "max_score", "order", "explanation"]
        widgets = {
            "text": forms.Textarea(attrs={"class": "rich"}),
            "explanation": forms.Textarea(attrs={"rows": 3}),
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["max_score"].required = False
        if self.instance.pk:
            options = list(self.instance.options.all())
            self.fields["options_text"].initial = "\n".join(o.text for o in options)
            for i, o in enumerate(options, start=1):
                if o.pk == self.instance.correct_option_id:
                    self.fields["correct_number"].initial = i
            self.fields["criteria_text"].initial = "\n".join(
                c.text if c.max_score == 2 else f"{c.text} | {c.max_score}" for c in self.instance.criteria.all()
            )

    def clean_text(self):
        value = clean_html(self.cleaned_data["text"])
        if is_empty_html(value):
            raise forms.ValidationError("Савол матнини ёзинг.")
        return value

    def clean(self):
        data = super().clean()
        options = [x.strip() for x in data.get("options_text", "").splitlines() if x.strip()]
        is_choice = data.get("qtype") == Question.Type.CHOICE
        if is_choice and len(options) < 2:
            self.add_error("options_text", "Камида 2 та вариант киритинг.")
        correct = data.get("correct_number")
        if correct and (not is_choice or correct > len(options)):
            self.add_error("correct_number", "Бундай рақамли вариант йўқ.")
        self.option_list = options if is_choice else []
        self.criteria_list = parse_criteria(data.get("criteria_text", ""))
        if self.criteria_list:
            data["max_score"] = Decimal(sum(s for _, s in self.criteria_list))
        elif correct and not data.get("max_score"):
            data["max_score"] = Decimal(1)
        if (data.get("max_score") or 0) <= 0:
            self.add_error("max_score", "Балл 0 дан катта бўлиши керак.")
        return data

    def save(self, commit=True):
        question = super().save(commit=False)
        question.max_score = self.cleaned_data["max_score"]
        question.correct_option = None
        question.save()
        question.options.all().delete()
        correct = self.cleaned_data.get("correct_number")
        for i, text in enumerate(self.option_list, start=1):
            option = question.options.create(text=text, order=i)
            if correct == i:
                question.correct_option = option
        question.save(update_fields=["correct_option"])
        question.criteria.all().delete()
        for i, (text, score) in enumerate(self.criteria_list, start=1):
            question.criteria.create(text=text, max_score=score, order=i)
        return question


class UserCreateForm(forms.ModelForm):
    password = forms.CharField(label="Парол", widget=forms.PasswordInput, min_length=8)

    class Meta:
        model = User
        fields = ["username", "last_name", "first_name", "role"]
        labels = {"username": "Логин", "last_name": "Фамилия", "first_name": "Исм"}
        help_texts = {"username": "Лотин ҳарфлари, рақамлар ва @/./+/-/_ белгилари."}

    def save(self, commit=True):
        user = super().save(commit=False)
        user.set_password(self.cleaned_data["password"])
        if commit:
            user.save()
        return user


class UserEditForm(forms.ModelForm):
    class Meta:
        model = User
        fields = ["username", "last_name", "first_name", "role", "is_active"]
        labels = {"username": "Логин", "last_name": "Фамилия", "first_name": "Исм", "is_active": "Фаол"}
        help_texts = {"username": "", "is_active": "Белгини олиб ташласангиз, фойдаланувчи тизимга кира олмайди."}


class UserPasswordForm(SetPasswordForm):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["new_password1"].label = "Янги парол"
        self.fields["new_password1"].help_text = "Камида 8 белги."
        self.fields["new_password2"].label = "Янги паролни такрорланг"
        self.error_messages["password_mismatch"] = "Пароллар бир хил эмас."


class ExtendTimeForm(forms.Form):
    minutes = forms.IntegerField(label="Қўшимча дақиқа", min_value=1, max_value=120, initial=10)
