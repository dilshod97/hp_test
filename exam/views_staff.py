"""Xodimlar sahifalari: test qo'shuvchi, tekshiruvchi va superadmin."""
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth import views as auth_views
from django.db import transaction
from django.db.models import Count, Max, Prefetch, Q
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from . import grading
from .decorators import role_required
from .forms import (
    BankForm, BlockForm, CandidateForm, CandidateImportForm, ExtendTimeForm, StaffLoginForm, QuestionForm, TestCheckersForm, TestForm,
    UserCreateForm, UserEditForm, UserPasswordForm,
)
from .models import (
    OPTION_LETTERS, Assignment, Attempt, Block, Candidate, Kind, Question, QuestionBank, Review, ReviewAnswer, Test, User,
)


class StaffLoginView(auth_views.LoginView):
    template_name = "exam/staff/login.html"
    authentication_form = StaffLoginForm
    redirect_authenticated_user = True


@role_required("is_superadmin", "is_author", "is_checker", "is_hr")
def dashboard(request):
    user = request.user
    if user.is_superadmin:
        return redirect("results")
    if user.is_hr:
        return redirect("candidate_list")
    if user.is_author:
        return redirect("test_list")
    return redirect("check_list")


def _locked_guard(request, test):
    if test.is_locked:
        messages.error(request, "Бу тестни номзодлар бошлаган — савол ва блокларни ўзгартириб бўлмайди.")
        return redirect("test_detail", test.pk)
    return None


def _question_guard(request, question):
    if question.block_id:
        return _locked_guard(request, question.block.test)
    return None


def _question_back(question):
    if question.block_id:
        return redirect("test_detail", question.block.test_id)
    return redirect("bank_detail", question.bank_id)


QUESTION_PREFETCH = ("options", "criteria")


# ---------------------------------------------------------------- Testlar

@role_required("is_author")
def test_list(request):
    tests = Test.objects.annotate(
        block_count=Count("blocks", distinct=True),
        assignment_count=Count("assignments", distinct=True),
    ).select_related("created_by")
    return render(request, "exam/staff/test_list.html", {"tests": tests})


@role_required("is_author")
def test_create(request):
    form = TestForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        test = form.save(commit=False)
        test.created_by = request.user
        test.save()
        form.save_m2m()
        messages.success(request, "Тест яратилди. Энди блок ва саволлар қўшинг.")
        return redirect("test_detail", test.pk)
    return render(request, "exam/staff/form.html", {"form": form, "title": "Янги тест"})


@role_required("is_author")
def test_detail(request, pk):
    test = get_object_or_404(Test, pk=pk)
    blocks = test.blocks.prefetch_related(
        Prefetch("questions", queryset=Question.objects.prefetch_related(*QUESTION_PREFETCH))
    )
    block_list = list(blocks)
    steps = [
        {"done": bool(block_list), "title": "Блок қўшинг", "hint": "Масалан, мантиқий тест ва ёзма амалий топшириқлар"},
        {"done": bool(block_list) and all(b.questions.all() for b in block_list), "title": "Ҳар бир блокка савол қўшинг",
         "hint": "«Базадан» тугмаси орқали тайёр саволларни танланг"},
        {"done": test.checkers.count() >= test.required_reviews, "title": "Экспертларни бириктиринг",
         "hint": f"Камида {test.required_reviews} та эксперт (суперадмин бириктиради)"},
        {"done": test.assignments.exists(), "title": "Номзодларни рўйхатга киритинг", "hint": "«Номзодлар» бўлимида"},
        {"done": test.is_active, "title": "Тестни фаол қилинг", "hint": "«Созламалар» → «Фаол»"},
    ]
    return render(request, "exam/staff/test_detail.html", {
        "test": test, "blocks": block_list, "steps": steps, "ready": all(x["done"] for x in steps),
        "interview_questions": test.interview_questions.order_by("bank__order", "order").prefetch_related("criteria"),
    })


@role_required("is_author")
def test_edit(request, pk):
    test = get_object_or_404(Test, pk=pk)
    form = TestForm(request.POST or None, instance=test)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Сақланди.")
        return redirect("test_detail", test.pk)
    return render(request, "exam/staff/form.html", {"form": form, "title": f"Тест созламалари: {test}", "back": test})


@role_required("is_author")
@require_POST
def test_delete(request, pk):
    test = get_object_or_404(Test, pk=pk)
    if test.is_locked:
        messages.error(request, "Натижалари бор тестни ўчириб бўлмайди. Уни нофаол қилинг.")
        return redirect("test_detail", test.pk)
    test.delete()
    messages.success(request, "Тест ўчирилди.")
    return redirect("test_list")


@role_required("is_author")
@require_POST
def test_copy(request, pk):
    """Qulflangan testni o'zgartirish uchun uning nusxasini yaratadi (nofaol, yangi kod bilan)."""
    src = get_object_or_404(Test, pk=pk)
    with transaction.atomic():
        test = Test.objects.create(
            title=f"{src.title} (нусха)", description=src.description, is_active=False,
            shuffle_questions=src.shuffle_questions, shuffle_options=src.shuffle_options,
            review_mode=src.review_mode, free_navigation=src.free_navigation, require_photo=src.require_photo, weight_practical=src.weight_practical, weight_logic=src.weight_logic,
            weight_interview=src.weight_interview, weight_defense=src.weight_defense, created_by=request.user,
        )
        test.checkers.set(src.checkers.all())
        test.interview_questions.set(src.interview_questions.all())
        for b in src.blocks.prefetch_related("questions__options", "questions__criteria"):
            nb = Block.objects.create(test=test, title=b.title, kind=b.kind, description=b.description,
                                      duration_minutes=b.duration_minutes, order=b.order)
            for q in b.questions.all():
                q.copy_to(block=nb)
    messages.success(request, "Нусха яратилди. Уни таҳрирлаб, кейин фаол қилинг.")
    return redirect("test_detail", test.pk)


@role_required("is_superadmin")
@require_POST
def test_unlock(request, pk):
    """Синов натижаларини ўчириб, тестни яна таҳрирлашга очади (пилот синовдан кейин)."""
    test = get_object_or_404(Test, pk=pk)
    n = test.attempts.count()
    with transaction.atomic():
        for attempt in test.attempts.all():
            attempt.delete()
    messages.success(request, f"{n} та натижа ўчирилди. Тестни яна таҳрирлаш мумкин.")
    return redirect("test_detail", test.pk)


@role_required("is_author")
def test_preview(request, pk):
    test = get_object_or_404(Test, pk=pk)
    blocks = test.blocks.prefetch_related("questions__options")
    return render(request, "exam/staff/test_preview.html", {"test": test, "blocks": blocks, "letters": OPTION_LETTERS})


BLOCK_PRESETS = {
    "logic": {
        "title": "Мантиқий фикрлаш тести", "kind": Kind.LOGIC, "duration_minutes": 40,
        "description": "Ҳар бир тестда битта жавобни белгиланг. Фақат шартдаги маълумотларга таянинг.",
    },
    "practical": {
        "title": "Ёзма амалий топшириқлар", "kind": Kind.PRACTICAL, "duration_minutes": 180,
        "description": "Ҳар бир топшириқни алоҳида бажаринг. Ҳисоб натижаси, ўлчов бирлиги ва фаразларингизни кўрсатинг.",
    },
}


@role_required("is_author")
def block_create(request, test_pk):
    test = get_object_or_404(Test, pk=test_pk)
    if guard := _locked_guard(request, test):
        return guard
    initial = {"order": test.blocks.count() + 1}
    initial.update(BLOCK_PRESETS.get(request.GET.get("preset"), {}))
    form = BlockForm(request.POST or None, initial=initial)
    if request.method == "POST" and form.is_valid():
        block = form.save(commit=False)
        block.test = test
        block.save()
        return redirect("test_detail", test.pk)
    return render(request, "exam/staff/form.html", {"form": form, "title": "Янги блок", "back": test})


@role_required("is_author")
def block_edit(request, pk):
    block = get_object_or_404(Block, pk=pk)
    if guard := _locked_guard(request, block.test):
        return guard
    form = BlockForm(request.POST or None, instance=block)
    if request.method == "POST" and form.is_valid():
        form.save()
        return redirect("test_detail", block.test_id)
    return render(request, "exam/staff/form.html", {"form": form, "title": f"Блок: {block}", "back": block.test})


@role_required("is_author")
@require_POST
def block_delete(request, pk):
    block = get_object_or_404(Block, pk=pk)
    if guard := _locked_guard(request, block.test):
        return guard
    block.delete()
    return redirect("test_detail", block.test_id)


@role_required("is_author")
def block_add_from_bank(request, pk):
    block = get_object_or_404(Block, pk=pk)
    if guard := _locked_guard(request, block.test):
        return guard
    banks = QuestionBank.objects.exclude(kind=Kind.INTERVIEW).annotate(n=Count("questions"))
    bank = banks.filter(pk=request.GET.get("bank") or request.POST.get("bank")).first()
    if bank is None:
        bank = banks.filter(kind=block.kind).first() or banks.first()

    if request.method == "POST":
        ids = [int(x) for x in request.POST.getlist("questions") if x.isdigit()]
        chosen = list(Question.objects.filter(pk__in=ids, bank__isnull=False).prefetch_related(*QUESTION_PREFETCH))
        chosen.sort(key=lambda q: (q.bank_id, q.order, q.pk))
        start = (block.questions.aggregate(m=Max("order"))["m"] or 0) + 1
        with transaction.atomic():
            for i, q in enumerate(chosen):
                q.copy_to(block=block, order=start + i)
        messages.success(request, f"{len(chosen)} та савол қўшилди.")
        return redirect("test_detail", block.test_id)

    already = set(block.questions.exclude(source=None).values_list("source_id", flat=True))
    questions = bank.questions.prefetch_related("criteria") if bank else []
    return render(request, "exam/staff/bank_pick.html", {
        "cur_block": block, "test": block.test, "banks": banks, "bank": bank,
        "questions": questions, "already": already,
    })


# ---------------------------------------------------------------- Savollar

@role_required("is_author")
def question_create(request, block_pk):
    block = get_object_or_404(Block, pk=block_pk)
    if guard := _locked_guard(request, block.test):
        return guard
    form = QuestionForm(request.POST or None, initial={"order": block.questions.count() + 1})
    form.instance.block = block
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            form.save()
        messages.success(request, "Савол қўшилди.")
        if "add_another" in request.POST:
            return redirect("question_create", block.pk)
        return redirect("test_detail", block.test_id)
    return render(request, "exam/staff/question_form.html", {
        "form": form, "place": block, "back_url": "test_detail", "back_obj": block.test,
    })


@role_required("is_author")
def bank_question_create(request, bank_pk):
    bank = get_object_or_404(QuestionBank, pk=bank_pk)
    form = QuestionForm(request.POST or None, initial={"order": bank.questions.count() + 1})
    form.instance.bank = bank
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            form.save()
        messages.success(request, "Савол базага қўшилди.")
        if "add_another" in request.POST:
            return redirect("bank_question_create", bank.pk)
        return redirect("bank_detail", bank.pk)
    return render(request, "exam/staff/question_form.html", {
        "form": form, "place": bank, "back_url": "bank_detail", "back_obj": bank,
    })


@role_required("is_author")
def question_edit(request, pk):
    q = get_object_or_404(Question.objects.select_related("block__test", "bank"), pk=pk)
    if guard := _question_guard(request, q):
        return guard
    form = QuestionForm(request.POST or None, instance=q)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            form.save()
        return _question_back(q)
    place = q.block or q.bank
    back_url, back_obj = ("test_detail", q.block.test) if q.block else ("bank_detail", q.bank)
    return render(request, "exam/staff/question_form.html", {
        "form": form, "place": place, "back_url": back_url, "back_obj": back_obj,
    })


@role_required("is_author")
@require_POST
def question_delete(request, pk):
    q = get_object_or_404(Question.objects.select_related("block__test"), pk=pk)
    if guard := _question_guard(request, q):
        return guard
    if q.answers.exists():
        messages.error(request, "Бу саволга жавоблар бор — ўчириб бўлмайди.")
        return _question_back(q)
    back = _question_back(q)
    q.delete()
    return back


# ---------------------------------------------------------------- Savollar banki

@role_required("is_author")
def bank_list(request):
    banks = QuestionBank.objects.annotate(n=Count("questions"))
    return render(request, "exam/staff/bank_list.html", {"banks": banks})


@role_required("is_author")
def bank_create(request):
    form = BankForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        bank = form.save()
        return redirect("bank_detail", bank.pk)
    return render(request, "exam/staff/form.html", {"form": form, "title": "Янги бўлим"})


@role_required("is_author")
def bank_edit(request, pk):
    bank = get_object_or_404(QuestionBank, pk=pk)
    form = BankForm(request.POST or None, instance=bank)
    if request.method == "POST" and form.is_valid():
        form.save()
        return redirect("bank_detail", bank.pk)
    return render(request, "exam/staff/form.html", {"form": form, "title": f"Бўлим: {bank}"})


@role_required("is_author")
def bank_detail(request, pk):
    bank = get_object_or_404(QuestionBank, pk=pk)
    questions = bank.questions.prefetch_related(*QUESTION_PREFETCH)
    return render(request, "exam/staff/bank_detail.html", {"bank": bank, "questions": questions})


# ---------------------------------------------------------------- Tekshiruvchi

@role_required("is_checker")
def check_list(request):
    Attempt.finish_abandoned()
    reviews = (
        Review.objects.filter(checker=request.user).select_related("attempt__test")
        .order_by("submitted_at", "assigned_at")
    )
    return render(request, "exam/staff/check_list.html", {"reviews": reviews})


def _review_rows(review, answers, items):
    rows = []
    for a in answers:
        scale = grading.answer_scale(a)
        item = items.get(a.pk)
        scores = list(item.scores) if item else []
        scores += [None] * (len(scale) - len(scores))
        rows.append({
            "answer": a,
            "criteria": [{"text": t, "max": m, "value": v, "choices": range(int(m) + 1)} for (t, m), v in zip(scale, scores)],
            "whole": not a.question.criteria.all(),
            "comment": item.comment if item else "",
            "total": item.total if item else None,
        })
    return rows


@role_required("is_checker")
def check_attempt(request, code):
    review = get_object_or_404(
        Review.objects.select_related("attempt__test"),
        attempt__code=code, checker=request.user,
    )
    attempt = review.attempt
    answers = list(
        grading.manual_answers(attempt).select_related("question__block")
        .prefetch_related("question__criteria", "question__options")
        .order_by("question__block__order", "question__block_id", "question__order", "question_id")
    )
    items = {i.answer_id: i for i in review.items.all()}

    if request.method == "POST":
        if review.is_submitted:
            messages.error(request, "Баҳолар топширилган — ўзгартириб бўлмайди.")
            return redirect("check_attempt", code=code)
        submit = "submit" in request.POST
        errors, missing = [], 0
        with transaction.atomic():
            for a in answers:
                scale = grading.answer_scale(a)
                scores = []
                for i, (_, mx) in enumerate(scale):
                    raw = request.POST.get(f"s_{a.pk}_{i}", "").strip().replace(",", ".")
                    if raw == "":
                        scores.append(None)
                        missing += 1
                        continue
                    try:
                        value = Decimal(raw)
                    except InvalidOperation:
                        value = Decimal(-1)
                    if value < 0 or value > Decimal(str(mx)):
                        errors.append(a.question.title or f"#{a.pk}")
                        scores.append(None)
                        missing += 1
                        continue
                    scores.append(int(value) if value == value.to_integral() else float(value))
                ReviewAnswer.objects.update_or_create(
                    review=review, answer=a,
                    defaults={"scores": scores, "comment": request.POST.get(f"c_{a.pk}", "").strip()},
                )
            if errors:
                messages.error(request, "Нотўғри балл: " + ", ".join(errors))
            elif submit and missing:
                messages.error(request, f"Топшириш учун барча мезонларни баҳоланг ({missing} та бўш).")
            elif submit:
                review.submitted_at = timezone.now()
                review.save(update_fields=["submitted_at"])
                grading.recompute(attempt)
                messages.success(request, "Баҳолар топширилди. Раҳмат!")
                return redirect("check_list")
            else:
                messages.success(request, "Қоралама сақланди.")
        return redirect("check_attempt", code=code)

    rows = _review_rows(review, answers, items)
    done = sum(1 for r in rows if r["total"] is not None)
    return render(request, "exam/staff/check_attempt.html", {
        "review": review, "attempt": attempt, "rows": rows,
        "total": sum((r["total"] for r in rows if r["total"] is not None), Decimal(0)),
        "done": done,
    })


# ---------------------------------------------------------------- Superadmin

@role_required("is_superadmin")
def user_list(request):
    users = User.objects.order_by("role", "username").annotate(
        pending=Count("reviews", filter=Q(reviews__submitted_at__isnull=True)),
    )
    return render(request, "exam/staff/user_list.html", {"users": users})


@role_required("is_superadmin")
def user_create(request):
    form = UserCreateForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Фойдаланувчи яратилди.")
        return redirect("user_list")
    return render(request, "exam/staff/form.html", {"form": form, "title": "Янги фойдаланувчи"})


@role_required("is_superadmin")
def user_edit(request, pk):
    user = get_object_or_404(User, pk=pk)
    form = UserEditForm(request.POST or None, instance=user, prefix="u")
    pw_form = UserPasswordForm(user, request.POST or None, prefix="p")
    if request.method == "POST":
        if "save_password" in request.POST:
            if pw_form.is_valid():
                pw_form.save()
                messages.success(request, "Парол ўзгартирилди.")
                return redirect("user_list")
        elif form.is_valid():
            if user == request.user and (not form.cleaned_data["is_active"] or form.cleaned_data["role"] != User.Role.SUPERADMIN):
                messages.error(request, "Ўзингизни блоклай ёки ролингизни ўзгартира олмайсиз.")
            else:
                form.save()
                messages.success(request, "Сақланди.")
                return redirect("user_list")
    return render(request, "exam/staff/user_edit.html", {"form": form, "pw_form": pw_form, "obj": user})


@role_required("is_superadmin")
def test_checkers(request, pk):
    test = get_object_or_404(Test, pk=pk)
    form = TestCheckersForm(request.POST or None, instance=test)
    if request.method == "POST" and form.is_valid():
        form.save()
        # Tekshiruvchisiz qolgan ishlarga endi biriktirish
        for attempt in test.attempts.filter(status=Attempt.Status.FINISHED):
            grading.ensure_reviews(attempt)
        messages.success(request, "Экспертлар бириктирилди.")
        return redirect("test_detail", test.pk)
    return render(request, "exam/staff/form.html", {"form": form, "title": f"Экспертлар: {test}", "back": test})


# ---------------------------------------------------------------- Номзодлар рўйхати (HR ва суперадмин)

def _register(last, first, middle, born, test):
    """born = (сана ёки None, йил). Қайтаради: (Assignment, янгими)."""
    born_date, year = born
    key = Candidate.make_identity_key(last, first, middle, born_date or year)
    candidate, _ = Candidate.objects.get_or_create(
        identity_key=key,
        defaults={"last_name": last.strip(), "first_name": first.strip(), "middle_name": middle.strip(),
                  "birth_date": born_date, "birth_year": year},
    )
    return Assignment.objects.get_or_create(candidate=candidate, test=test)


@role_required("can_manage_candidates")
def candidate_list(request):
    qs = Assignment.objects.select_related("candidate", "test").order_by("-created_at")
    test_id = request.GET.get("test")
    search = request.GET.get("q", "").strip()
    if test_id:
        qs = qs.filter(test_id=test_id)
    if search:
        qs = qs.filter(Q(candidate__last_name__icontains=search) | Q(candidate__first_name__icontains=search)
                       | Q(login=search))
    attempts = {(a.candidate_id, a.test_id): a for a in Attempt.objects.filter(
        candidate_id__in=qs.values("candidate_id"))}
    rows = [{"x": x, "attempt": attempts.get((x.candidate_id, x.test_id))} for x in qs]
    return render(request, "exam/staff/candidate_list.html", {
        "rows": rows, "tests": Test.objects.all(), "params": request.GET,
        "form": CandidateForm(initial={"test": test_id}), "import_form": CandidateImportForm(initial={"test": test_id}),
        "new_id": request.GET.get("new"),
    })


@role_required("can_manage_candidates")
@require_POST
def candidate_add(request):
    form = CandidateForm(request.POST)
    if form.is_valid():
        d = form.cleaned_data
        assignment, created = _register(d["last_name"], d["first_name"], d["middle_name"], d["birth"], d["test"])
        if created:
            messages.success(request, f"{d['last_name']} {d['first_name']} рўйхатга қўшилди. "
                                      f"Логин: {assignment.login}, парол: {assignment.password}")
        else:
            messages.info(request, "Бу номзод шу тест рўйхатида аллақачон бор.")
        return redirect(f"{reverse('candidate_list')}?test={d['test'].pk}&new={assignment.pk}")
    messages.error(request, "Маълумотлар нотўғри: " + "; ".join(e for errs in form.errors.values() for e in errs))
    return redirect(f"{reverse('candidate_list')}?test={request.POST.get('test', '')}")


@role_required("can_manage_candidates")
@require_POST
def candidate_import(request):
    form = CandidateImportForm(request.POST, request.FILES)
    if form.is_valid():
        test = form.cleaned_data["test"]
        with transaction.atomic():
            added = sum(_register(*row, test)[1] for row in form.rows)
        messages.success(request, f"Файлдан {len(form.rows)} та номзод ўқилди, {added} та янги қўшилди. "
                                  f"Логин ва паролларни «Барчасини чоп этиш» орқали чиқаринг.")
    else:
        messages.error(request, " ".join(e for errs in form.errors.values() for e in errs))
    return redirect(f"{reverse('candidate_list')}?test={request.POST.get('test', '')}")


@role_required("can_manage_candidates")
@require_POST
def assignment_delete(request, pk):
    x = get_object_or_404(Assignment, pk=pk)
    if x.attempt:
        messages.error(request, "Номзод тестни бошлаган — рўйхатдан ўчириб бўлмайди.")
    else:
        x.delete()
        messages.success(request, "Рўйхатдан ўчирилди.")
    return redirect(f"{reverse('candidate_list')}?test={x.test_id}")


@role_required("can_manage_candidates")
@require_POST
def assignment_reset_password(request, pk):
    x = get_object_or_404(Assignment, pk=pk)
    x.reset_password()
    messages.success(request, f"{x.candidate.full_name}: янги парол {x.password}. Эски варақча энди ишламайди.")
    return redirect(f"{reverse('candidate_list')}?test={x.test_id}&new={x.pk}")


@role_required("can_manage_candidates")
def candidate_print(request):
    """Номзодларга бериладиган варақчалар (логин ва парол)."""
    qs = Assignment.objects.select_related("candidate", "test").order_by("candidate__last_name", "candidate__first_name")
    ids = [int(x) for x in request.GET.getlist("id") if x.isdigit()]
    if ids:
        qs = qs.filter(pk__in=ids)
    elif request.GET.get("test"):
        qs = qs.filter(test_id=request.GET["test"])
    else:
        qs = qs.none()
    done = set(Attempt.objects.exclude(status=Attempt.Status.IN_PROGRESS).values_list("candidate_id", "test_id"))
    slips = [x for x in qs if (x.candidate_id, x.test_id) not in done]
    exam_url = request.build_absolute_uri(reverse("exam_login"))
    return render(request, "exam/staff/candidate_print.html", {"slips": slips, "exam_url": exam_url,
                                                               "skipped": len(qs) - len(slips) if not ids else 0})


@role_required("can_manage_candidates")
def candidate_template(request):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Номзодлар"
    ws.append(["Фамилия", "Исм", "Отасининг исми", "Туғилган йили"])
    ws.append(["Тошматов", "Бобур", "Акмалович", "1993"])
    for col, w in zip("ABCD", [18, 14, 20, 15]):
        ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="nomzodlar_namuna.xlsx"'
    return response


def _results_queryset(request):
    qs = Attempt.objects.select_related("candidate", "test").prefetch_related("reviews__checker")
    test_id = request.GET.get("test")
    status = request.GET.get("status")
    search = request.GET.get("q", "").strip()
    if test_id:
        qs = qs.filter(test_id=test_id)
    if status:
        qs = qs.filter(status=status)
    if search:
        qs = qs.filter(
            Q(candidate__last_name__icontains=search) | Q(candidate__first_name__icontains=search)
            | Q(code__icontains=search)
        )
    return qs.order_by("-started_at")


@role_required("is_superadmin")
def results(request):
    Attempt.finish_abandoned()
    Attempt.purge_old_photos()
    rows = [{"a": a, "s": grading.summary(a)} for a in _results_queryset(request)]
    base = Attempt.objects.all()
    if request.GET.get("test"):
        base = base.filter(test_id=request.GET["test"])
    counts = {r["status"]: r["n"] for r in base.values("status").annotate(n=Count("id"))}
    registered = Assignment.objects.filter(**({"test_id": request.GET["test"]} if request.GET.get("test") else {})).count()
    return render(request, "exam/staff/results.html", {
        "rows": rows,
        "stats": {
            "registered": registered,
            "in_progress": counts.get(Attempt.Status.IN_PROGRESS, 0),
            "finished": counts.get(Attempt.Status.FINISHED, 0),
            "graded": counts.get(Attempt.Status.GRADED, 0),
        },
        "tests": Test.objects.all(),
        "statuses": Attempt.Status.choices,
        "params": request.GET,
    })


@role_required("is_superadmin")
def result_detail(request, pk):
    attempt = get_object_or_404(Attempt.objects.select_related("candidate", "test"), pk=pk)
    reviews = list(attempt.reviews.select_related("checker").order_by("round", "id"))
    items = {}
    for item in ReviewAnswer.objects.filter(review__attempt=attempt, review__submitted_at__isnull=False):
        items[(item.review_id, item.answer_id)] = item
    answers = list(
        attempt.answers.select_related("question__block", "selected_option")
        .prefetch_related("question__options", "question__criteria")
        .order_by("question__block__order", "question__block_id", "question__order", "question_id")
    )
    rows = []
    for a in answers:
        scale = grading.answer_scale(a)
        per_review = []
        for r in reviews:
            item = items.get((r.pk, a.pk))
            per_review.append({"review": r, "item": item})
        criteria_rows = []
        for i, (text, mx) in enumerate(scale):
            vals = []
            for pr in per_review:
                s = pr["item"].scores if pr["item"] else []
                vals.append(s[i] if i < len(s) else None)
            criteria_rows.append({"text": text, "max": mx, "values": vals})
        letters = {o.pk: OPTION_LETTERS[i] for i, o in enumerate(a.question.options.all())}
        rows.append({
            "answer": a, "per_review": per_review, "criteria": criteria_rows,
            "chosen": letters.get(a.selected_option_id, ""), "key": letters.get(a.question.correct_option_id, ""),
        })

    interview = []
    for q in attempt.test.interview_questions.order_by("bank__order", "order").prefetch_related("criteria"):
        vals = attempt.interview_scores.get(str(q.pk), [])
        crit = list(q.criteria.all())
        interview.append({
            "q": q,
            "criteria": [{"c": c, "i": i, "value": vals[i] if i < len(vals) else None, "choices": range(c.max_score + 1)}
                         for i, c in enumerate(crit)],
        })

    reviewer_ids = [r.checker_id for r in reviews]
    return render(request, "exam/staff/result_detail.html", {
        "attempt": attempt, "rows": rows, "reviews": reviews,
        "summary": grading.summary(attempt),
        "block_attempts": attempt.block_attempts.select_related("block"),
        "free_checkers": User.objects.filter(role=User.Role.CHECKER, is_active=True).exclude(pk__in=reviewer_ids),
        "interview": interview,
        "letters": OPTION_LETTERS,
    })


@role_required("is_superadmin")
@require_POST
def result_action(request, pk):
    attempt = get_object_or_404(Attempt, pk=pk)
    action = request.POST.get("action")
    if action == "add_review":
        checker = User.objects.filter(pk=request.POST.get("checker"), role=User.Role.CHECKER).first()
        if checker and not attempt.reviews.filter(checker=checker).exists():
            Review.objects.create(attempt=attempt, checker=checker, round=attempt.reviews.count() + 1)
            grading.recompute(attempt)
            messages.success(request, f"{checker} эксперт сифатида қўшилди.")
    elif action in ("reopen_review", "delete_review"):
        review = get_object_or_404(Review, pk=request.POST.get("review"), attempt=attempt)
        if action == "reopen_review":
            review.submitted_at = None
            review.save(update_fields=["submitted_at"])
            messages.success(request, "Текширув қайта очилди — эксперт баҳоларни ўзгартира олади.")
        elif review.is_submitted:
            messages.error(request, "Топширилган текширувни ўчириб бўлмайди. Аввал қайта очинг.")
        else:
            review.delete()
            messages.success(request, "Текширув олиб ташланди.")
        grading.recompute(attempt)
    elif action == "interview":
        scores = {}
        for q in attempt.test.interview_questions.prefetch_related("criteria"):
            vals = []
            for i, c in enumerate(q.criteria.all()):
                raw = request.POST.get(f"iv_{q.pk}_{i}", "")
                vals.append(int(raw) if raw.isdigit() and int(raw) <= c.max_score else None)
            scores[str(q.pk)] = vals
        attempt.interview_scores = scores
        attempt.interview_comment = request.POST.get("interview_comment", "").strip()
        raw = request.POST.get("defense_score", "").strip().replace(",", ".")
        try:
            defense = Decimal(raw) if raw else None
        except InvalidOperation:
            defense = Decimal(-1)
        if defense is not None and not (0 <= defense <= attempt.test.weight_defense):
            messages.error(request, f"Ҳимоя бали 0 дан {attempt.test.weight_defense:g} гача бўлиши керак.")
        else:
            attempt.defense_score = defense
            messages.success(request, "Суҳбат ва ҳимоя баллари сақланди.")
        attempt.save(update_fields=["interview_scores", "interview_comment", "defense_score"])
    elif action == "finish" and attempt.status == Attempt.Status.IN_PROGRESS:
        attempt.finish()
        messages.success(request, "Тест мажбуран якунланди.")
    elif action == "extend" and attempt.status == Attempt.Status.IN_PROGRESS:
        form = ExtendTimeForm(request.POST)
        ba = attempt.current_block_attempt()
        if form.is_valid() and ba:
            # Эркин режимда барча блоклар битта умумий вақтга эга
            targets = attempt.block_attempts.filter(finished_at__isnull=True) if attempt.test.free_navigation else [ba]
            for target in targets:
                target.extra_minutes += form.cleaned_data["minutes"]
                target.save(update_fields=["extra_minutes"])
            messages.success(request, f"Жорий блокка {form.cleaned_data['minutes']} дақиқа қўшилди.")
    elif action == "delete":
        attempt.delete()
        messages.success(request, "Натижа ўчирилди. Номзод тестни қайтадан бошлаши мумкин.")
        return redirect("results")
    return redirect("result_detail", attempt.pk)


def _num(value):
    return float(value) if value is not None else None


@role_required("is_superadmin")
def results_export(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Натижалар"
    headers = [
        "Код", "Фамилия", "Исм", "Отасининг исми", "Туғилган сана", "Тест", "Ҳолат",
        "Расм", "Мантиқий тест", "Мантиқий макс.", "Талқин", "Ёзма амалий", "Амалий макс.",
        "Суҳбат", "Суҳбат макс.", "Ҳимоя", "Якуний (100)", "Учинчи текширув", "Экспертлар",
        "Бошлаган", "Топширган",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    tz = timezone.get_current_timezone()
    for a in _results_queryset(request):
        s = grading.summary(a)
        logic, practical, interview = s["logic"] or {}, s["practical"] or {}, s["interview"] or {}
        ws.append([
            a.code, a.candidate.last_name, a.candidate.first_name, a.candidate.middle_name,
            a.candidate.birth_date or a.candidate.birth_year, a.test.title, a.get_status_display(),
            a.get_photo_status_display() if a.photo_status else ("Олинмаган" if a.test.require_photo else "—"),
            _num(logic.get("score")), _num(logic.get("max")), s["logic_band"],
            _num(practical.get("score")), _num(practical.get("max")),
            _num(interview.get("score")), _num(interview.get("max")),
            _num(s["defense"]), _num(s["final"]),
            a.third_review_reason if a.third_review_needed else "",
            ", ".join(str(r.checker) for r in a.reviews.all()),
            a.started_at.astimezone(tz).replace(tzinfo=None),
            a.finished_at.astimezone(tz).replace(tzinfo=None) if a.finished_at else None,
        ])
    widths = [10, 16, 14, 18, 14, 30, 13, 16, 10, 10, 20, 10, 10, 9, 9, 9, 11, 28, 30, 18, 18]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="natijalar_{timezone.localdate():%Y-%m-%d}.xlsx"'
    return response


# ---------------------------------------------------------------- Расм юклаш (савол матни учун)

IMAGE_SIGNATURES = {
    b"\x89PNG\r\n\x1a\n": ".png", b"\xff\xd8\xff": ".jpg", b"GIF87a": ".gif", b"GIF89a": ".gif",
}


@role_required("is_author")
@require_POST
def upload_image(request):
    import uuid

    from django.conf import settings
    from django.http import JsonResponse

    f = request.FILES.get("file")
    if f is None:
        return JsonResponse({"error": "Файл йўқ"}, status=400)
    if f.size > settings.IMAGE_UPLOAD_MAX_BYTES:
        return JsonResponse({"error": "Расм 5 МБ дан катта бўлмасин"}, status=400)
    head = f.read(16)
    ext = next((e for sig, e in IMAGE_SIGNATURES.items() if head.startswith(sig)), None)
    if ext is None and head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        ext = ".webp"
    if ext is None:
        return JsonResponse({"error": "Фақат PNG, JPG, GIF ёки WEBP расм"}, status=400)
    f.seek(0)
    folder = settings.MEDIA_ROOT / "questions"
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{uuid.uuid4().hex}{ext}"
    with open(folder / name, "wb") as out:
        for chunk in f.chunks():
            out.write(chunk)
    return JsonResponse({"location": f"{settings.MEDIA_URL}questions/{name}"})


@role_required("is_superadmin")
def result_photo(request, pk):
    """Номзод расми — фақат суперадмин учун (экспертларга ҳеч қачон берилмайди)."""
    from django.http import FileResponse

    attempt = get_object_or_404(Attempt, pk=pk)
    path = attempt.photo_path
    if not path or not path.exists():
        raise Http404
    response = FileResponse(open(path, "rb"), content_type="image/jpeg")
    response["Cache-Control"] = "private, no-store"
    return response


# ---------------------------------------------------------------- Ҳужжатдан юклаш (.docx)

@role_required("is_author")
def bank_import(request):
    """1-қадам: файлни юклаб, текшириш. 2-қадам: тасдиқлаб, базага ёзиш."""
    import uuid

    from django.conf import settings

    from .importer import import_parsed, parse

    folder = settings.MEDIA_ROOT / "imports"
    folder.mkdir(parents=True, exist_ok=True)
    ctx = {"retention": settings.PHOTO_RETENTION_DAYS}

    if request.method == "POST" and request.POST.get("step") == "import":
        token = request.POST.get("token", "")
        path = folder / f"{token}.docx"
        if not token.isalnum() or not path.exists():
            messages.error(request, "Файл топилмади, қайтадан юкланг.")
            return redirect("bank_import")
        parsed = parse(str(path))
        if parsed.errors:
            messages.error(request, "Ҳужжатда хатолар бор.")
            return redirect("bank_import")
        source = request.POST.get("source", "").strip()[:255] or parsed.title or "Ҳужжат"
        report = import_parsed(parsed, source, user=request.user, create_tests=bool(request.POST.get("create_tests")))
        path.unlink(missing_ok=True)
        text = f"Юкланди: {report['created']} та янги, {report['updated']} та янгиланган савол."
        if report["tests"]:
            text += " Намуна тестлар (нофаол): " + ", ".join(report["tests"]) + "."
        messages.success(request, text)
        return redirect("bank_list")

    if request.method == "POST":
        f = request.FILES.get("file")
        if f is None or not f.name.lower().endswith(".docx"):
            messages.error(request, "Word (.docx) файлини танланг.")
            return redirect("bank_import")
        if f.size > 20 * 1024 * 1024:
            messages.error(request, "Файл 20 МБ дан катта бўлмасин.")
            return redirect("bank_import")
        # Эски (тасдиқланмаган) юкламаларни тозалаймиз
        import time
        for old in folder.glob("*.docx"):
            if time.time() - old.stat().st_mtime > 24 * 3600:
                old.unlink(missing_ok=True)
        token = uuid.uuid4().hex
        path = folder / f"{token}.docx"
        with open(path, "wb") as out:
            for chunk in f.chunks():
                out.write(chunk)
        parsed = parse(str(path))
        source = parsed.title or f.name.rsplit(".", 1)[0]
        existing = QuestionBank.objects.filter(source=source).count()
        ctx.update({"parsed": parsed, "token": token, "source": source, "filename": f.name, "existing": existing})
        if parsed.errors:
            path.unlink(missing_ok=True)
    return render(request, "exam/staff/bank_import.html", ctx)
