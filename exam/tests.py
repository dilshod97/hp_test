import hashlib
import json
import os
from datetime import timedelta
from decimal import Decimal

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from . import grading
from .models import Assignment, Attempt, Candidate, Kind, Question, QuestionBank, Review, Test, User

DOCX = os.environ.get(
    "TANLOV_DOCX", os.path.expanduser("~/Downloads/Hisob_palatasi_tanlov_materiallari_uz_kirill.docx")
)


class Base(TestCase):
    review_mode = Test.ReviewMode.DOUBLE

    def setUp(self):
        from django.core.cache import cache
        cache.clear()  # кириш уринишлари ҳисоблагичи тестлар орасида тўпланмасин
        self.admin = User.objects.create_user("admin", password="x" * 10, role=User.Role.SUPERADMIN)
        self.author = User.objects.create_user("author", password="x" * 10, role=User.Role.AUTHOR)
        self.c1 = User.objects.create_user("c1", password="x" * 10, role=User.Role.CHECKER)
        self.c2 = User.objects.create_user("c2", password="x" * 10, role=User.Role.CHECKER)
        self.c3 = User.objects.create_user("c3", password="x" * 10, role=User.Role.CHECKER)

        # Bank: suhbat savoli
        iv_bank = QuestionBank.objects.create(name="Suhbat", kind=Kind.INTERVIEW)
        self.iv = Question.objects.create(bank=iv_bank, code="С 01", title="Bosim", text="<p>?</p>", qtype="text", max_score=10)
        for i in range(5):
            self.iv.criteria.create(text=f"m{i}", max_score=2, order=i)

        self.client.force_login(self.author)
        self.client.post(reverse("test_create"), {
            "title": "Auditor", "is_active": "on", "shuffle_questions": "on", "review_mode": self.review_mode,
            "interview_questions": [self.iv.pk],
            "weight_practical": 55, "weight_logic": 15, "weight_interview": 15, "weight_defense": 15,
        })
        self.test = Test.objects.get()
        self.client.post(reverse("block_create", args=[self.test.pk]),
                         {"title": "Mantiq", "kind": "logic", "duration_minutes": 40, "order": 1})
        self.client.post(reverse("block_create", args=[self.test.pk]),
                         {"title": "Amaliy", "kind": "practical", "duration_minutes": 180, "order": 2})
        self.logic, self.practical = self.test.blocks.all()

        r = self.client.post(reverse("question_create", args=[self.logic.pk]), {
            "text": "<p>2+2?</p><script>alert(1)</script>", "qtype": "choice", "order": 1,
            "options_text": "3\n4\n5\n6\n7", "correct_number": 2,
        })
        self.assertEqual(r.status_code, 302, r.content[:2000])
        self.client.post(reverse("question_create", args=[self.logic.pk]), {
            "text": "<p>3+3?</p>", "qtype": "choice", "order": 2, "options_text": "5\n6\n7", "correct_number": 2,
        })
        for code in ("МА 01", "МА 02"):
            r = self.client.post(reverse("question_create", args=[self.practical.pk]), {
                "title": f"Vazifa {code}", "text": f"<p>{code} sharti</p>", "qtype": "text", "order": 1,
                "criteria_text": "a\nb\nc\nd\ne",
            })
            self.assertEqual(r.status_code, 302)
        self.test.checkers.set([self.c1, self.c2, self.c3])
        self.client.logout()

        self.q_logic1, self.q_logic2 = self.logic.questions.all()
        self.q_ma1, self.q_ma2 = self.practical.questions.all()
        self.assignment = self.register("Karimov", "Aziz", "Olimovich", "1995")

    def register(self, last, first, middle, born, test=None):
        self.client.force_login(self.admin)
        self.client.post(reverse("candidate_add"), {
            "last_name": last, "first_name": first, "middle_name": middle, "birth": born,
            "test": (test or self.test).pk,
        })
        self.client.logout()
        return Assignment.objects.get(candidate__last_name=last, candidate__first_name=first, test=test or self.test)

    def login_candidate(self, client=None, login=None, password=None, assignment=None):
        a = assignment or self.assignment
        data = {"login": login or a.login, "password": password or a.password}
        return (client or self.client).post(reverse("exam_login"), data)

    def save(self, payload):
        return self.client.post(reverse("exam_save"), json.dumps(payload), content_type="application/json")

    def take_test(self, logic_correct=True):
        self.login_candidate()
        self.client.get(reverse("exam_take"))
        opt = self.q_logic1.options.get(text="4" if logic_correct else "3")
        self.assertTrue(self.save({"question": self.q_logic1.pk, "option": opt.pk}).json()["ok"])
        attempt = Attempt.objects.get()
        self.client.post(reverse("exam_next"), {"block_attempt": attempt.block_attempts.first().pk})
        self.client.get(reverse("exam_take"))
        self.assertTrue(self.save({"question": self.q_ma1.pk, "text": "<p>Javob 475</p>"}).json()["ok"])
        self.assertTrue(self.save({"question": self.q_ma2.pk, "text": "<p>Ikkinchi</p>"}).json()["ok"])
        self.client.post(reverse("exam_next"), {"block_attempt": attempt.block_attempts.last().pk})
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.FINISHED)
        self.client.get(reverse("exam_done"))
        return attempt

    def grade(self, checker, attempt, scores_by_question, submit=True):
        """scores_by_question: {question: [m1..m5]}"""
        self.client.force_login(checker)
        post = {"submit": "1"} if submit else {"save": "1"}
        for q, scores in scores_by_question.items():
            a = attempt.answers.get(question=q)
            for i, s in enumerate(scores):
                post[f"s_{a.pk}_{i}"] = "" if s is None else str(s)
            post[f"c_{a.pk}"] = f"izoh-{checker.username}"
        return self.client.post(reverse("check_attempt", args=[attempt.code]), post)


class FlowTests(Base):
    def test_question_form_parses_key_and_criteria(self):
        self.assertNotIn("<script", self.q_logic1.text)
        self.assertEqual(self.q_logic1.correct_option.text, "4")
        self.assertEqual(self.q_logic1.max_score, 1)
        self.assertTrue(self.q_logic1.is_auto)
        self.assertEqual(self.q_ma1.criteria.count(), 5)
        self.assertEqual(self.q_ma1.max_score, 10)
        self.assertEqual(self.test.max_score_for(Kind.PRACTICAL), 20)

    def test_candidate_sees_letters_not_key_or_criteria(self):
        self.login_candidate()
        r = self.client.get(reverse("exam_take"))
        self.assertContains(r, 'class="letter">А<')
        self.assertContains(r, 'class="letter">Д<')
        self.assertNotContains(r, "to&#x27;g&#x27;ri")
        self.assertNotContains(r, "kalit")

    def test_double_review_average(self):
        attempt = self.take_test()
        # Mantiqiy: 1 to'g'ri, 1 bo'sh -> avtomatik
        self.assertEqual(attempt.answers.get(question=self.q_logic1).score, 1)
        self.assertEqual(attempt.answers.get(question=self.q_logic2).score, 0)
        reviews = list(attempt.reviews.all())
        self.assertEqual(len(reviews), 2)
        first, second = [r.checker for r in reviews]

        # Ekspert ism ko'rmaydi
        self.client.force_login(first)
        r = self.client.get(reverse("check_attempt", args=[attempt.code]))
        self.assertContains(r, "Javob 475")
        self.assertNotContains(r, "<td>М1")
        self.assertNotContains(r, "Karimov")
        self.assertNotContains(r, "1995")
        self.assertNotContains(r, "2+2")  # avtomatik savollar ko'rsatilmaydi

        # Topshirish uchun hamma mezon kerak
        r = self.grade(first, attempt, {self.q_ma1: [2, 2, 2, 2, None], self.q_ma2: [1] * 5}, submit=True)
        self.assertFalse(Review.objects.get(attempt=attempt, checker=first).is_submitted)

        self.grade(first, attempt, {self.q_ma1: [2, 2, 2, 2, 1], self.q_ma2: [1] * 5})
        self.assertTrue(Review.objects.get(attempt=attempt, checker=first).is_submitted)
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.FINISHED)

        # Topshirilgandan keyin o'zgartirib bo'lmaydi
        self.grade(first, attempt, {self.q_ma1: [0] * 5, self.q_ma2: [0] * 5})
        self.assertEqual(Review.objects.get(attempt=attempt, checker=first).items.get(answer__question=self.q_ma1).scores, [2, 2, 2, 2, 1])

        # Ikkinchi ekspert birinchisining balini ko'rmaydi
        self.client.force_login(second)
        r = self.client.get(reverse("check_attempt", args=[attempt.code]))
        self.assertNotContains(r, f"izoh-{first.username}")
        self.assertNotContains(r, 'selected>2')
        self.grade(second, attempt, {self.q_ma1: [2, 1, 2, 2, 2], self.q_ma2: [1, 1, 1, 2, 1]})

        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.GRADED)
        self.assertFalse(attempt.third_review_needed)
        # MA1: (9 + 9) / 2 = 9; MA2: (5 + 6) / 2 = 5.5
        self.assertEqual(attempt.answers.get(question=self.q_ma1).score, Decimal("9"))
        self.assertEqual(attempt.answers.get(question=self.q_ma2).score, Decimal("5.5"))

        # Yakuniy: amaliy 55*14.5/20 = 39.88, mantiqiy 15*1/2 = 7.5; suhbat va himoya yo'q -> final None
        s = grading.summary(attempt)
        self.assertEqual(s["practical"]["weighted"], Decimal("39.88"))
        self.assertEqual(s["logic"]["weighted"], Decimal("7.50"))
        self.assertIsNone(s["final"])

        # Superadmin suhbat va himoyani kiritadi
        self.client.force_login(self.admin)
        post = {"action": "interview", "defense_score": "12"}
        for i, v in enumerate([2, 2, 1, 2, 1]):
            post[f"iv_{self.iv.pk}_{i}"] = str(v)
        self.client.post(reverse("result_action", args=[attempt.pk]), post)
        attempt.refresh_from_db()
        s = grading.summary(attempt)
        self.assertEqual(s["interview"]["score"], 8)
        self.assertEqual(s["final"], Decimal("39.88") + Decimal("7.50") + Decimal("12.00") + Decimal("12"))

        r = self.client.get(reverse("results"))
        self.assertContains(r, "Karimov Aziz Olimovich")
        r = self.client.get(reverse("result_detail", args=[attempt.pk]))
        self.assertContains(r, "Эксп. 2")
        r = self.client.get(reverse("results_export"))
        self.assertTrue(r.content.startswith(b"PK"))

    def test_third_review_on_large_difference_uses_median(self):
        attempt = self.take_test()
        first, second = [r.checker for r in attempt.reviews.all()]
        self.grade(first, attempt, {self.q_ma1: [2, 2, 2, 2, 2], self.q_ma2: [2, 2, 2, 1, 1]})
        self.grade(second, attempt, {self.q_ma1: [1, 1, 1, 1, 1], self.q_ma2: [1, 1, 1, 1, 1]})
        attempt.refresh_from_db()
        # Farq 18 - 10 = 8 > 20 * 15% = 3
        self.assertTrue(attempt.third_review_needed)
        self.assertIn("фарқи", attempt.third_review_reason)
        self.assertEqual(attempt.status, Attempt.Status.FINISHED)
        third = attempt.reviews.get(round=3)
        self.assertNotIn(third.checker, (first, second))
        self.grade(third.checker, attempt, {self.q_ma1: [1, 2, 1, 2, 1], self.q_ma2: [1, 2, 2, 2, 1]})
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.GRADED)
        # MA1 mediana: [2,1,1]=1, [2,1,2]=2, [2,1,1]=1, [2,1,2]=2, [2,1,1]=1 -> 7
        self.assertEqual(attempt.answers.get(question=self.q_ma1).score, 7)
        # MA2: [2,1,1]=1, [2,1,2]=2, [2,1,2]=2, [1,1,2]=1, [1,1,1]=1 -> 7
        self.assertEqual(attempt.answers.get(question=self.q_ma2).score, 7)

    def test_third_review_on_zero_vs_two(self):
        attempt = self.take_test()
        first, second = [r.checker for r in attempt.reviews.all()]
        self.grade(first, attempt, {self.q_ma1: [2, 2, 2, 2, 0], self.q_ma2: [1] * 5})
        self.grade(second, attempt, {self.q_ma1: [2, 2, 2, 2, 2], self.q_ma2: [1] * 5})
        attempt.refresh_from_db()
        self.assertTrue(attempt.third_review_needed)
        self.assertIn("Vazifa МА 01", attempt.third_review_reason)

    def test_superadmin_reopen_and_add_reviewer(self):
        attempt = self.take_test()
        first, second = [r.checker for r in attempt.reviews.all()]
        self.grade(first, attempt, {self.q_ma1: [2] * 5, self.q_ma2: [2] * 5})
        review = attempt.reviews.get(checker=first)
        self.client.force_login(self.admin)
        self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "reopen_review", "review": review.pk})
        review.refresh_from_db()
        self.assertFalse(review.is_submitted)
        r2 = attempt.reviews.get(checker=second)
        self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "delete_review", "review": r2.pk})
        third = (set([self.c1, self.c2, self.c3]) - {first, second}).pop()
        self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "add_review", "checker": third.pk})
        self.assertTrue(attempt.reviews.filter(checker=third).exists())

    def test_checker_isolation_and_roles(self):
        attempt = self.take_test()
        reviewers = {r.checker for r in attempt.reviews.all()}
        outsider = ({self.c1, self.c2, self.c3} - reviewers).pop()
        self.client.force_login(outsider)
        self.assertEqual(self.client.get(reverse("check_attempt", args=[attempt.code])).status_code, 404)
        self.assertEqual(self.client.get(reverse("results")).status_code, 403)
        self.assertEqual(self.client.get(reverse("test_list")).status_code, 403)
        self.assertEqual(self.client.get(reverse("bank_list")).status_code, 403)
        self.client.force_login(self.author)
        self.assertEqual(self.client.get(reverse("user_list")).status_code, 403)
        self.assertEqual(self.client.get(reverse("check_list")).status_code, 403)

    def test_duplicate_blocked_and_lock(self):
        self.take_test()
        r = self.login_candidate()
        self.assertContains(r, "суперадмин рухсати")
        self.client.force_login(self.author)
        r = self.client.get(reverse("question_edit", args=[self.q_ma1.pk]))
        self.assertRedirects(r, reverse("test_detail", args=[self.test.pk]))

    def test_block_timer_expires_on_server(self):
        self.login_candidate()
        self.client.get(reverse("exam_take"))
        attempt = Attempt.objects.get()
        ba1 = attempt.block_attempts.first()
        ba1.started_at = timezone.now() - timedelta(minutes=41)
        ba1.save()
        opt = self.q_logic1.options.first()
        self.assertEqual(self.save({"question": self.q_logic1.pk, "option": opt.pk}).status_code, 409)
        r = self.client.get(reverse("exam_take"))
        self.assertContains(r, "Amaliy")

    def test_resume_after_restart(self):
        self.login_candidate()
        self.client.get(reverse("exam_take"))
        r = self.login_candidate(client=self.client_class())
        self.assertRedirects(r, reverse("exam_take"))
        self.assertEqual(Attempt.objects.count(), 1)

    def test_weights_must_sum_to_100(self):
        self.client.force_login(self.author)
        r = self.client.post(reverse("test_edit", args=[self.test.pk]), {
            "title": "X", "review_mode": "single",
            "weight_practical": 50, "weight_logic": 15, "weight_interview": 15, "weight_defense": 15,
        })
        self.assertContains(r, "100 бўлиши керак")

    def test_copy_test(self):
        self.take_test()
        self.client.force_login(self.author)
        self.client.post(reverse("test_copy", args=[self.test.pk]))
        copy = Test.objects.exclude(pk=self.test.pk).get()
        self.assertFalse(copy.is_active)
        self.assertEqual(copy.max_score, self.test.max_score)
        q = Question.objects.get(block__test=copy, title="Vazifa МА 01")
        self.assertEqual(q.criteria.count(), 5)
        q = Question.objects.filter(block__test=copy, qtype="choice").first()
        self.assertEqual(q.correct_option.question, q)
        self.assertEqual([x.pk for x in copy.interview_questions.all()], [self.iv.pk])

    def test_bank_pick(self):
        bank = QuestionBank.objects.create(name="MA", kind=Kind.PRACTICAL)
        src = Question.objects.create(bank=bank, code="МА 07", title="Kredit", text="<p>t</p>", qtype="text", max_score=10)
        src.criteria.create(text="x", max_score=2)
        self.client.force_login(self.author)
        r = self.client.get(reverse("block_add_from_bank", args=[self.practical.pk]) + f"?bank={bank.pk}")
        self.assertContains(r, "Kredit")
        self.client.post(reverse("block_add_from_bank", args=[self.practical.pk]), {"bank": bank.pk, "questions": [src.pk]})
        copy = self.practical.questions.get(title="Kredit")
        self.assertEqual(copy.source, src)
        self.assertEqual(copy.criteria.count(), 1)
        # Bankdagi o'zgarish nusxaga ta'sir qilmaydi
        src.title = "Boshqa"
        src.save()
        copy.refresh_from_db()
        self.assertEqual(copy.title, "Kredit")


class SingleReviewTests(Base):
    review_mode = Test.ReviewMode.SINGLE

    def test_single_review(self):
        attempt = self.take_test(logic_correct=False)
        self.assertEqual(attempt.reviews.count(), 1)
        self.assertEqual(attempt.answers.get(question=self.q_logic1).score, 0)
        checker = attempt.reviews.get().checker
        self.grade(checker, attempt, {self.q_ma1: [2, 2, 2, 2, 2], self.q_ma2: [0, 0, 0, 0, 0]})
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.GRADED)
        self.assertEqual(attempt.total_score, 10)


class SebTests(Base):
    @override_settings(REQUIRE_SEB=True)
    def test_seb_required(self):
        self.assertEqual(self.client.get(reverse("exam_login")).status_code, 403)
        r = self.client.get(reverse("exam_login"), HTTP_USER_AGENT="Mozilla/5.0 SEB/3.7")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.client.get(reverse("staff_login")).status_code, 200)

    @override_settings(SEB_CONFIG_KEYS=["abc123"])
    def test_seb_config_key(self):
        url = "http://testserver" + reverse("exam_login")
        good = hashlib.sha256((url + "abc123").encode()).hexdigest()
        self.assertEqual(self.client.get(reverse("exam_login")).status_code, 403)
        self.assertEqual(self.client.get(reverse("exam_login"), HTTP_X_SAFEEXAMBROWSER_CONFIGKEYHASH=good).status_code, 200)


class ImportTests(TestCase):
    def setUp(self):
        if not os.path.exists(DOCX):
            self.skipTest("Tanlov materiallari fayli topilmadi")
        User.objects.create_user("admin", password="x" * 10, role=User.Role.SUPERADMIN)

    def test_import(self):
        call_command("import_tanlov", DOCX, verbosity=0)
        self.assertEqual(Question.objects.filter(bank__isnull=False).count(), 220)
        self.assertEqual(Question.objects.filter(bank__kind=Kind.LOGIC, correct_option__isnull=False).count(), 100)
        self.assertEqual(QuestionBank.objects.filter(kind=Kind.LOGIC).count(), 1)  # 5 тўплам битта бўлимда
        self.assertEqual(Test.objects.count(), 2)
        for t in Test.objects.all():
            self.assertEqual(t.max_score_for(Kind.LOGIC), 20)
            self.assertEqual(t.max_score_for(Kind.PRACTICAL), 60)
            self.assertEqual(t.interview_questions.count(), 5)
            self.assertEqual(t.total_minutes, 220)
        # Kodlar (МА 01, Т1-01) ekranda ko'rinmaydi
        admin = User.objects.get(username="admin")
        self.client.force_login(admin)
        t = Test.objects.get(title__startswith="Молиявий")
        r = self.client.get(reverse("test_detail", args=[t.pk]))
        self.assertContains(r, "Банк қолдиғини солиштириш")
        self.assertNotContains(r, "МА 01")
        self.assertNotContains(r, "Т1-01")
        # Qayta import — dublikat yo'q
        call_command("import_tanlov", DOCX, verbosity=0)
        self.assertEqual(Question.objects.filter(bank__isnull=False).count(), 220)
        self.assertEqual(Test.objects.count(), 2)


class MergedTaskWarningTests(TestCase):
    def test_restarted_criteria_numbering_warns(self):
        from .importer import Parsed, Section, _validate
        crit = lambda n: [(f"мезон {i}", 2) for i in range(1, n + 1)]
        merged = {"code": "МА 32", "title": "", "text": ["шарт"], "criteria": crit(5) * 2,
                  "crit_nums": [str(i) for i in range(1, 6)] * 2}
        normal = {"code": "МА 43", "title": "", "text": ["шарт"], "criteria": crit(4),
                  "crit_nums": ["1", "2", "3", "4"]}
        out = Parsed(sections=[Section("practical", "МА", items=[merged, normal])])
        _validate(out)
        self.assertEqual(len(out.warnings), 1)
        self.assertIn("МА 32", out.warnings[0])


class LanguageTests(Base):
    """Номзод киргач тилни танлайди: ўзбекча — асосий тест, русча — унинг рус варианти."""

    def setUp(self):
        super().setUp()
        self.client.force_login(self.author)
        self.client.post(reverse("test_copy", args=[self.test.pk]))
        self.client.logout()
        self.ru = Test.objects.exclude(pk=self.test.pk).get()
        self.ru.title, self.ru.language, self.ru.is_active = "Аудитор (рус)", "ru", True
        self.ru.save()
        self.test.ru_version = self.ru
        self.test.save()

    def test_candidate_chooses_russian(self):
        r = self.login_candidate()
        self.assertRedirects(r, reverse("exam_lang"))
        r = self.client.get(reverse("exam_lang"))
        self.assertContains(r, "Русский")
        self.assertContains(r, "Аудитор (рус)")
        r = self.client.post(reverse("exam_lang"), {"lang": "ru"})
        self.assertRedirects(r, reverse("exam_take"))
        attempt = Attempt.objects.get()
        self.assertEqual(attempt.test, self.ru)
        self.assertEqual(self.assignment.attempt, attempt)
        r = self.client.get(reverse("exam_take"))
        self.assertContains(r, "Завершить блок")  # блоклар кетма-кет режимда
        self.assertContains(r, 'lang: "ru"')
        self.assertNotContains(r, "Блокни якунлаш")
        # Ҳамма матнлар таржима қилинган: ўзбек ҳарфлари (ў, қ, ғ, ҳ) кўринадиган матнда қолмаган
        import re
        visible = re.sub(r"<script.*?</script>|<[^>]+>", " ", r.content.decode(), flags=re.S)
        self.assertFalse(re.search(r"[ЎўҚқҒғҲҳ]", visible), re.findall(r".{30}[ЎўҚқҒғҲҳ].{30}", visible)[:3])
        # Рўйхатда номзоднинг ҳолати асосий тест бўйича кўринади
        self.client.force_login(self.admin)
        r = self.client.get(reverse("candidate_list"))
        self.assertContains(r, "Ечилмоқда")
        self.assertNotContains(r, "Бошламаган")

    def test_choose_uzbek_and_no_second_attempt(self):
        self.login_candidate()
        self.client.post(reverse("exam_lang"), {"lang": "uz"})
        attempt = Attempt.objects.get()
        self.assertEqual(attempt.test, self.test)
        attempt.finish()
        other = self.client_class()
        r = self.login_candidate(client=other)
        self.assertContains(r, "топширгансиз")
        self.assertEqual(Attempt.objects.count(), 1)

    def test_inactive_russian_version_skips_choice(self):
        self.ru.is_active = False
        self.ru.save()
        r = self.login_candidate()
        self.assertRedirects(r, reverse("exam_take"), fetch_redirect_response=False)
        self.assertEqual(Attempt.objects.get().test, self.test)

    def test_choice_page_requires_login(self):
        r = self.client.post(reverse("exam_lang"), {"lang": "ru"})
        self.assertRedirects(r, reverse("exam_login"))
        self.assertFalse(Attempt.objects.exists())

    def test_login_page_switches_language(self):
        self.client.get(reverse("exam_login") + "?lang=ru")
        r = self.client.get(reverse("exam_login"))
        self.assertContains(r, "Вход в тест")
        self.assertContains(r, "Пароль")

    def test_hr_cannot_assign_russian_version(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("candidate_add"), {
            "last_name": "Petrov", "first_name": "Ivan", "middle_name": "", "birth": "1990", "test": self.ru.pk})
        self.assertFalse(Assignment.objects.filter(test=self.ru).exists())

    def test_russian_document_import(self):
        path = os.path.expanduser("~/Downloads/Schetnaya_palata_materialy_otbora_ru_редакт.docx")
        if not os.path.exists(path):
            self.skipTest("Рус тилидаги ҳужжат топилмади")
        from .importer import import_parsed, parse
        parsed = parse(path)
        self.assertEqual((parsed.lang, parsed.errors, parsed.warnings), ("ru", [], []))
        import_parsed(parsed, source="ru.docx", user=self.admin)
        banks = QuestionBank.objects.filter(source="ru.docx")
        self.assertEqual(banks.get(kind=Kind.LOGIC).name, "Логический тест")
        self.assertEqual(banks.get(kind=Kind.LOGIC).questions.filter(correct_option__isnull=False).count(), 100)
        fa = banks.get(name__contains="финансового аудитора", kind=Kind.PRACTICAL)
        self.assertEqual(fa.questions.count(), 48)
        self.assertFalse(fa.questions.filter(code__in=["ФА 33", "ФА 34"]).exists())
        self.assertEqual(fa.questions.get(code="ФА 43").criteria.count(), 4)
        self.assertEqual(banks.filter(kind=Kind.INTERVIEW).count(), 3)


class McqImportTests(TestCase):
    """Тўғри жавоби сариқ белгиланган вариантли тестлар (иқтисод, ACCA)."""

    def make_doc(self):
        import tempfile

        import docx
        from docx.enum.text import WD_COLOR_INDEX
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        d = docx.Document()
        d.add_heading("ACCA бўйича тест саволлари", 1)
        d.add_paragraph("Битта тўғри жавоб вариантини танланг.")
        def opt(text, mark=None):
            p = d.add_paragraph()
            r = p.add_run(text)
            if mark == "hl":
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if mark == "shd":
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "FFF200")
                p._p.get_or_add_pPr().append(shd)
            return p
        opt("1. Биринчи савол?"); opt("А) бир"); opt("Б) икки", "hl"); opt("В) уч"); opt("Г) тўрт")
        d.add_paragraph("Иккинчи савол (рўйхат)?", style="List Paragraph")
        opt("A) йўқ", "shd"); opt("B) ҳа"); opt("C) балки")
        opt("3. Учинчи савол?"); opt("- 5%"); opt("+ 4,6%"); opt("- 9%")
        opt("4. Ҳисобланг: 10 + 5"); opt("Жавоб:"); opt("10 млн сўм"); opt("15 млн сўм", "hl")
        p = opt("5. Амортизация — пул жамғариш."); p = d.add_paragraph(); p.add_run("Тўғри / "); r = p.add_run("Нотўғри")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        f = tempfile.NamedTemporaryFile(suffix=".docx", delete=False); d.save(f.name)
        return f.name

    def test_parse_all_formats(self):
        from .mcq_importer import import_mcq, parse_mcq
        parsed = parse_mcq(self.make_doc())
        self.assertEqual(parsed.errors, [])
        got = [(q["text"], q["options"][q["correct"][0]]) for q in parsed.items]
        self.assertEqual(got, [("Биринчи савол?", "икки"), ("Иккинчи савол (рўйхат)?", "йўқ"),
                               ("Учинчи савол?", "4,6%"), ("Ҳисобланг: 10 + 5", "15 млн сўм"),
                               ("Амортизация — пул жамғариш.", "Нотўғри")])
        r = import_mcq(parsed, bank_name="ACCA", source="t", prefix="ACCA")
        self.assertEqual((r["created"], r["bank"].kind), (5, Kind.TEST))
        r = import_mcq(parsed, bank_name="ACCA", source="t", prefix="ACCA")   # такрор — дубликатсиз
        self.assertEqual((r["created"], r["updated"], r["bank"].questions.count()), (0, 5, 5))

    def test_unmarked_question_is_an_error(self):
        import tempfile

        import docx
        from .mcq_importer import parse_mcq
        d = docx.Document()
        for t in ("1. Савол?", "А) бир", "Б) икки"):
            d.add_paragraph(t)
        f = tempfile.NamedTemporaryFile(suffix=".docx", delete=False); d.save(f.name)
        self.assertIn("тўғри жавоб 0 та", parse_mcq(f.name).errors[0])


class ProfessionalTestScoringTests(Base):
    def test_hundred_points_with_professional_block(self):
        """Тестлар 2 баллдан, ёзма топшириқ 6 баллдан: якуний балл 100 лик."""
        self.client.force_login(self.author)
        self.client.post(reverse("block_create", args=[self.test.pk]),
                         {"title": "ACCA", "kind": "test", "duration_minutes": 20, "order": 3})
        acca = self.test.blocks.get(kind=Kind.TEST)
        self.client.post(reverse("question_create", args=[acca.pk]), {
            "text": "<p>ACCA?</p>", "qtype": "choice", "order": 1, "options_text": "a\nb", "correct_number": 1})
        # 3 та тест (2 мантиқий + 1 касбий) × 2 = 6 балл → вазн 6 ва 2 × 47 = 94 ёзма
        Test.objects.filter(pk=self.test.pk).update(weight_logic=4, weight_test=2, weight_practical=94,
                                                    weight_interview=0, weight_defense=0)
        self.client.logout()
        self.login_candidate()
        attempt = Attempt.objects.get()
        q = acca.questions.get()
        attempt.answers.filter(question=q).update(selected_option=q.correct_option)  # блоклар кетма-кет — тўғридан
        attempt.finish()
        s = grading.summary(attempt)
        self.assertEqual((s["test"]["score"], s["test"]["max"]), (1, 1))
        self.assertEqual(s["logic"]["score"], 0)
        # Устунларда 100 баллик шкаладаги бал (2/2), остида хом натижа (1 / 1)
        self.assertEqual((s["test"]["weighted"], s["test"]["weight"]), (2, 2))
        self.client.force_login(self.admin)
        for url in (reverse("results"), reverse("result_detail", args=[attempt.pk])):
            r = self.client.get(url)
            self.assertContains(r, "Касбий")
            self.assertContains(r, "/ 2")
            self.assertNotContains(r, "Суҳбат</th>")   # вазни 0 — устун яширин
        self.assertEqual(self.client.get(reverse("results_export")).status_code, 200)
        # Тартиблаш: жадвал sortable, катакларда локалсиз сон
        r = self.client.get(reverse("results"))
        self.assertContains(r, 'class="sortable"')
        self.assertContains(r, 'data-v="2.00"')
        # Тест саҳифаси: хом балл йиғиндиси эмас, якуний 100 балл ва ҳар саволнинг улуши
        r = self.client.get(reverse("test_detail", args=[self.test.pk]))
        self.assertContains(r, "Якуний балл: <b>100</b> (мантиқий 4 + касбий 2 + ёзма 94)")
        self.assertContains(r, "ҳар бир савол 2 балл")
        self.assertContains(r, "ҳар бир савол 47 балл")


class SessionTests(Base):
    def test_candidate_login_drops_staff_session(self):
        self.client.force_login(self.admin)
        self.login_candidate()
        self.assertEqual(self.client.get(reverse("results")).status_code, 302)
        self.assertEqual(self.client.get(reverse("exam_take")).status_code, 200)


class RegistrationTests(Base):
    def test_credentials_generated(self):
        a = self.assignment
        self.assertRegex(a.login, r"^\d{6}$")
        self.assertRegex(a.password, r"^[2-9]{6}$")
        self.assertEqual(a.candidate.birth_year, 1995)
        self.assertIsNone(a.candidate.birth_date)

    def test_wrong_password_and_lock(self):
        r = self.login_candidate(password="000000")
        self.assertContains(r, "Логин ёки парол нотўғри")
        r = self.login_candidate(login="999999")
        self.assertContains(r, "Логин ёки парол нотўғри")
        for _ in range(5):
            self.login_candidate(password="000000")
        r = self.login_candidate()  # тўғри парол ҳам блокланган
        self.assertContains(r, "кўп марта нотўғри")
        self.assertEqual(Attempt.objects.count(), 0)

    def test_inactive_test_rejected(self):
        self.test.is_active = False
        self.test.save()
        r = self.login_candidate()
        self.assertContains(r, "ҳали очилмаган")

    def test_each_test_has_own_login(self):
        other = Test.objects.create(title="Tahlilchi", is_active=True, created_by=self.admin, require_photo=False)
        other.blocks.create(title="B", kind="logic", duration_minutes=5)
        a2 = self.register("Karimov", "Aziz", "Olimovich", "1995", test=other)
        self.assertNotEqual(a2.login, self.assignment.login)
        r = self.login_candidate(assignment=a2)
        self.assertRedirects(r, reverse("exam_take"))
        self.assertEqual(Attempt.objects.get().test, other)

    def test_retake_needs_superadmin(self):
        attempt = self.take_test()
        r = self.login_candidate()
        self.assertContains(r, "суперадмин рухсати")
        hr = User.objects.create_user("hr", password="x" * 10, role=User.Role.HR)
        self.client.force_login(hr)
        self.assertEqual(self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "delete"}).status_code, 403)
        self.client.force_login(self.admin)
        self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "delete"})
        self.client.logout()
        self.assertRedirects(self.login_candidate(), reverse("exam_take"))

    def test_hr_role(self):
        hr = User.objects.create_user("hr", password="x" * 10, role=User.Role.HR)
        self.client.force_login(hr)
        self.assertRedirects(self.client.get(reverse("dashboard")), reverse("candidate_list"))
        r = self.client.get(reverse("candidate_list"))
        self.assertContains(r, self.assignment.login)
        self.assertContains(r, self.assignment.password)
        for name in ("results", "test_list", "bank_list", "user_list", "check_list"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)
        self.client.post(reverse("candidate_add"), {"last_name": "Ali", "first_name": "Vali", "middle_name": "",
                                                    "birth": "12.04.1994", "test": self.test.pk})
        new = Assignment.objects.get(candidate__last_name="Ali")
        self.assertEqual(new.candidate.birth_year, 1994)
        r = self.client.get(reverse("candidate_print") + f"?test={self.test.pk}")
        self.assertContains(r, new.login)
        self.assertContains(r, new.password)
        self.assertContains(r, "Ali Vali")
        old = new.password
        self.client.post(reverse("assignment_reset_password", args=[new.pk]))
        new.refresh_from_db()
        self.assertNotEqual(new.password, old)

    def test_print_skips_finished(self):
        self.take_test()
        self.client.force_login(self.admin)
        r = self.client.get(reverse("candidate_print") + f"?test={self.test.pk}")
        self.assertNotContains(r, self.assignment.password)
        r = self.client.get(reverse("candidate_print") + f"?id={self.assignment.pk}")
        self.assertContains(r, "тестни топшириб бўлган")

    def test_excel_import(self):
        from io import BytesIO

        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Фамилия", "Исм", "Отасининг исми", "Туғилган йили"])
        ws.append(["Тошматов", "Бобур", "Акмалович", 1993])
        ws.append(["Каримова", "Нилуфар", "", "11.02.1996"])
        ws.append([None, None, None, None])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "list.xlsx"
        self.client.force_login(self.admin)
        self.client.post(reverse("candidate_import"), {"file": buf, "test": self.test.pk})
        self.assertEqual(Assignment.objects.filter(test=self.test).count(), 3)
        self.client.logout()
        a = Assignment.objects.get(candidate__last_name="Тошматов")
        self.assertEqual(a.candidate.birth_year, 1993)
        self.assertRedirects(self.login_candidate(assignment=a), reverse("exam_take"))

    def test_assignment_delete_only_before_start(self):
        x = Assignment.objects.get()
        self.login_candidate()
        self.client.force_login(self.admin)
        self.client.post(reverse("assignment_delete", args=[x.pk]))
        self.assertTrue(Assignment.objects.filter(pk=x.pk).exists())


class EditorContentTests(Base):
    def test_sanitizers(self):
        from .utils import clean_answer, clean_html
        html = ('<h3 style="text-align:center;position:fixed">S</h3><p style="color:#b91c1c">r</p>'
                '<pre>x = 1</pre><img src="/media/questions/a.png"><img src="http://evil/x.png">'
                '<a href="http://x">l</a><table border="1"><tr><td style="width:50%">c</td></tr></table>')
        author = clean_html(html)
        self.assertIn('text-align:center', author)
        self.assertNotIn('position', author)
        self.assertIn('color:#b91c1c', author)
        self.assertIn('<pre>', author)
        self.assertIn('src="/media/questions/a.png"', author)
        self.assertNotIn('evil', author)
        self.assertNotIn('<a', author)
        answer = clean_answer(html)
        self.assertNotIn('<img', answer)
        self.assertIn('<table', answer)

    def test_image_upload(self):
        import tempfile
        from django.core.files.uploadedfile import SimpleUploadedFile
        png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
        with tempfile.TemporaryDirectory() as tmp, self.settings(MEDIA_ROOT=__import__("pathlib").Path(tmp)):
            self.client.force_login(self.c1)
            r = self.client.post(reverse("upload_image"), {"file": SimpleUploadedFile("a.png", png)})
            self.assertEqual(r.status_code, 403)
            self.client.force_login(self.author)
            r = self.client.post(reverse("upload_image"), {"file": SimpleUploadedFile("a.png", png)})
            self.assertEqual(r.status_code, 200)
            self.assertTrue(r.json()["location"].startswith("/media/questions/"))
            r = self.client.post(reverse("upload_image"), {"file": SimpleUploadedFile("a.png", b"<script>")})
            self.assertEqual(r.status_code, 400)


class TemplateVarTests(Base):
    def test_bank_pick_breadcrumb(self):
        self.client.force_login(self.author)
        r = self.client.get(reverse("block_add_from_bank", args=[self.practical.pk]))
        self.assertContains(r, "Amaliy")
        self.assertNotContains(r, "BlockNode")


class FreeNavigationTests(Base):
    def setUp(self):
        super().setUp()
        self.test.free_navigation = True
        self.test.save()

    def test_all_blocks_open_and_editable(self):
        self.login_candidate()
        r = self.client.get(reverse("exam_take"))
        # Иккала блок саволлари бир саҳифада
        self.assertContains(r, "2+2?")
        self.assertContains(r, "МА 02 sharti")
        attempt = Attempt.objects.get()
        deadlines = {ba.deadline for ba in attempt.block_attempts.all()}
        self.assertEqual(len(deadlines), 1)  # битта умумий муддат
        total = (deadlines.pop() - attempt.block_attempts.first().started_at).total_seconds() / 60
        self.assertAlmostEqual(total, 220, delta=0.1)
        # Аввал амалий, кейин мантиқий саволга қайтиб жавоб ўзгартириш
        self.assertTrue(self.save({"question": self.q_ma2.pk, "text": "<p>javob</p>"}).json()["ok"])
        opt = self.q_logic1.options.get(text="4")
        self.assertTrue(self.save({"question": self.q_logic1.pk, "option": opt.pk}).json()["ok"])
        opt = self.q_logic1.options.get(text="5")
        self.assertTrue(self.save({"question": self.q_logic1.pk, "option": opt.pk}).json()["ok"])
        # Якунлаш бутун тестни топширади
        self.client.post(reverse("exam_next"), {"block_attempt": attempt.block_attempts.first().pk})
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.FINISHED)
        self.assertEqual(attempt.answers.get(question=self.q_logic1).score, 0)

    def test_total_time_expires(self):
        self.login_candidate()
        self.client.get(reverse("exam_take"))
        attempt = Attempt.objects.get()
        attempt.block_attempts.update(started_at=timezone.now() - timedelta(minutes=221))
        self.client.get(reverse("exam_take"))
        attempt.refresh_from_db()
        self.assertNotEqual(attempt.status, Attempt.Status.IN_PROGRESS)
        self.assertFalse(attempt.block_attempts.filter(finished_at__isnull=True).exists())

    def test_extend_all_blocks(self):
        self.login_candidate()
        self.client.get(reverse("exam_take"))
        attempt = Attempt.objects.get()
        self.client.force_login(self.admin)
        self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "extend", "minutes": 10})
        self.assertEqual(len({ba.deadline for ba in attempt.block_attempts.all()}), 1)


class SetupAndUnlockTests(Base):
    def test_checklist_and_presets(self):
        self.client.force_login(self.author)
        t = Test.objects.create(title="Yangi", created_by=self.author)
        r = self.client.get(reverse("test_detail", args=[t.pk]))
        self.assertContains(r, "Тестни тайёрлаш")
        r = self.client.get(reverse("block_create", args=[t.pk]) + "?preset=logic")
        self.assertContains(r, "Мантиқий фикрлаш тести")
        r = self.client.post(reverse("block_create", args=[t.pk]), {"title": "", "kind": "logic", "duration_minutes": 40, "order": 1})
        self.assertContains(r, "Бу майдонни тўлдиринг")

    def test_unlock(self):
        self.login_candidate()
        self.client.get(reverse("exam_take"))
        self.assertTrue(self.test.is_locked)
        self.client.force_login(self.author)
        self.assertEqual(self.client.post(reverse("test_unlock", args=[self.test.pk])).status_code, 403)
        self.client.force_login(self.admin)
        self.client.post(reverse("test_unlock", args=[self.test.pk]))
        self.assertFalse(self.test.is_locked)
        self.assertTrue(Assignment.objects.filter(test=self.test).exists())  # рўйхат сақланади


class PhotoTests(Base):
    JPEG = "data:image/jpeg;base64," + __import__("base64").b64encode(b"\xff\xd8\xff\xe0" + b"\x00" * 64).decode()

    def setUp(self):
        super().setUp()
        import tempfile
        from pathlib import Path
        self.tmp = tempfile.TemporaryDirectory()
        self.media = self.settings(MEDIA_ROOT=Path(self.tmp.name))
        self.media.enable()
        self.test.require_photo = True
        self.test.save()

    def tearDown(self):
        self.media.disable()
        self.tmp.cleanup()

    def post_photo(self, payload):
        return self.client.post(reverse("exam_photo_save"), json.dumps(payload), content_type="application/json")

    def test_photo_flow_and_privacy(self):
        self.login_candidate()
        self.assertRedirects(self.client.get(reverse("exam_take")), reverse("exam_photo"))
        self.assertContains(self.client.get(reverse("exam_photo")), "Шахсни тасдиқлаш")
        self.assertEqual(self.post_photo({"image": "data:image/jpeg;base64,PHNjcmlwdD4="}).status_code, 400)
        self.assertTrue(self.post_photo({"image": self.JPEG}).json()["ok"])
        attempt = Attempt.objects.get()
        self.assertEqual(attempt.photo_status, "taken")
        self.assertTrue(attempt.photo_path.exists())
        self.assertEqual(self.client.get(reverse("exam_take")).status_code, 200)

        # Тестни якунлаб, экспертга юборамиз
        ba1, ba2 = attempt.block_attempts.all()
        self.client.post(reverse("exam_next"), {"block_attempt": ba1.pk})
        self.client.get(reverse("exam_take"))
        self.assertTrue(self.save({"question": self.q_ma1.pk, "text": "<p>javob</p>"}).json()["ok"])
        self.client.post(reverse("exam_next"), {"block_attempt": ba2.pk})
        attempt.refresh_from_db()
        url = reverse("result_photo", args=[attempt.pk])
        checker = attempt.reviews.first().checker
        self.client.force_login(checker)
        self.assertNotContains(self.client.get(reverse("check_attempt", args=[attempt.code])), url)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertContains(self.client.get(reverse("result_detail", args=[attempt.pk])), url)

        path = attempt.photo_path
        attempt.delete()
        self.assertFalse(path.exists())

    def test_no_camera_continues(self):
        self.login_candidate()
        self.assertTrue(self.post_photo({"status": "no_camera"}).json()["ok"])
        self.assertEqual(self.client.get(reverse("exam_take")).status_code, 200)
        self.assertEqual(Attempt.objects.get().photo_status, "no_camera")

    def test_purge(self):
        self.login_candidate()
        self.post_photo({"image": self.JPEG})
        attempt = Attempt.objects.get()
        Attempt.objects.filter(pk=attempt.pk).update(photo_at=timezone.now() - timedelta(days=181))
        self.assertEqual(Attempt.purge_old_photos(), 1)
        attempt.refresh_from_db()
        self.assertEqual((attempt.photo, attempt.photo_status), ("", "purged"))


def make_docx(title="Ички аудит танлови", logic_key="Б", code="ИТ"):
    from io import BytesIO

    import docx
    d = docx.Document()
    d.add_paragraph(title, style="Title")
    d.add_paragraph("Ички аудитор учун амалий топшириқлар", style="Heading 1")
    d.add_paragraph("Кириш матни.")
    d.add_paragraph(f"{code} 01 Тармоқ хавфсизлиги", style="Heading 2")
    d.add_paragraph("Тармоқдаги хавфни баҳоланг.")
    d.add_paragraph("1. Хавфни аниқлайди. 2 балл")
    d.add_paragraph("2. Чора таклиф этади. 3 балл")
    d.add_paragraph("1 тўплам", style="Heading 1")
    d.add_paragraph("1. 2+2 нечага тенг?")
    d.add_paragraph("А) 3     Б) 4     В) 5")
    d.add_paragraph("Тест жавоблари ва изоҳлари", style="Heading 1")
    d.add_paragraph("1 тўплам калити", style="Heading 2")
    t = d.add_table(rows=2, cols=4)
    for i, v in enumerate(["№", "Тўғри", "Яқин вариантлар", "Қисқа изоҳ"]):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(["1", logic_key, "А, В", "Оддий қўшиш"]):
        t.rows[1].cells[i].text = v
    buf = BytesIO()
    d.save(buf)
    buf.seek(0)
    buf.name = "hujjat.docx"
    return buf


class WebImportTests(Base):
    def setUp(self):
        super().setUp()
        import tempfile
        from pathlib import Path
        self.tmp = tempfile.TemporaryDirectory()
        self.media = self.settings(MEDIA_ROOT=Path(self.tmp.name))
        self.media.enable()
        self.client.force_login(self.author)

    def tearDown(self):
        self.media.disable()
        self.tmp.cleanup()

    def upload(self, buf, source=None):
        r = self.client.post(reverse("bank_import"), {"step": "check", "file": buf})
        self.assertEqual(r.status_code, 200)
        token = r.context.get("token")
        if r.context["parsed"].errors:
            return r, None
        r2 = self.client.post(reverse("bank_import"), {"step": "import", "token": token,
                                                       "source": source or r.context["source"]})
        return r, r2

    def test_check_then_import(self):
        r, r2 = self.upload(make_docx())
        self.assertContains(r, "Юклашга тайёр")
        self.assertContains(r, "Ички аудитор учун амалий топшириқлар")
        self.assertRedirects(r2, reverse("bank_list"))
        banks = QuestionBank.objects.filter(source="Ички аудит танлови")
        self.assertEqual(banks.count(), 2)
        task = Question.objects.get(bank__source="Ички аудит танлови", code="ИТ 01")
        self.assertEqual(task.max_score, 5)
        self.assertEqual(task.criteria.count(), 2)
        logic = Question.objects.get(bank__source="Ички аудит танлови", qtype="choice")
        self.assertEqual(logic.correct_option.text, "4")
        # Қайта юклаш — дубликат йўқ, калит янгиланади
        self.upload(make_docx(logic_key="В"))
        self.assertEqual(Question.objects.filter(bank__source="Ички аудит танлови").count(), 2)
        logic = Question.objects.get(bank__source="Ички аудит танлови", qtype="choice")
        self.assertEqual(logic.correct_option.text, "5")
        # Бошқа ҳужжат бошқа манба остида — мавжудлар устига ёзмайди
        self.upload(make_docx(title="Бошқа танлов"))
        self.assertEqual(Question.objects.filter(bank__source="Ички аудит танлови").count(), 2)
        self.assertEqual(Question.objects.filter(bank__source="Бошқа танлов").count(), 2)

    def test_errors_block_import(self):
        r, r2 = self.upload(make_docx(logic_key="Е"))
        self.assertIsNone(r2)
        self.assertContains(r, "калитдаги")
        self.assertEqual(QuestionBank.objects.filter(source="Ички аудит танлови").count(), 0)
        from io import BytesIO
        bad = BytesIO(b"not a docx"); bad.name = "x.docx"
        r = self.client.post(reverse("bank_import"), {"step": "check", "file": bad})
        self.assertContains(r, "Файлни ўқиб бўлмади")

    def test_experts_cannot_import(self):
        self.client.force_login(self.c1)
        self.assertEqual(self.client.get(reverse("bank_import")).status_code, 403)


class ConfirmSubmitTests(Base):
    def test_confirm_uses_prototype_submit(self):
        # Баҳолаш формасида name="submit" тугма бор — form.submit() уни кўрмайди, прототип орқали юбориш шарт
        self.client.force_login(self.author)
        r = self.client.get(reverse("test_list"))
        self.assertContains(r, "HTMLFormElement.prototype.submit.call(form)")


class RandomPickTests(Base):
    def test_random_n_from_bank(self):
        bank = QuestionBank.objects.create(name="Katta bank", kind=Kind.LOGIC)
        for i in range(30):
            q = Question.objects.create(bank=bank, code=f"R{i}", title=f"Savol {i}", text=f"<p>{i}</p>", qtype="choice", max_score=1)
            o = q.options.create(text="ha", order=1); q.options.create(text="yo'q", order=2)
            q.correct_option = o; q.save()
        self.client.force_login(self.author)
        url = reverse("block_add_from_bank", args=[self.logic.pk])
        before = self.logic.questions.count()
        self.client.post(url, {"bank": bank.pk, "mode": "random", "count": "10"})
        added = self.logic.questions.filter(source__bank=bank)
        self.assertEqual(added.count(), 10)
        self.assertEqual(self.logic.questions.count(), before + 10)
        self.assertTrue(all(q.correct_option and q.correct_option.question_id == q.pk for q in added))
        # Иккинчи марта — блокдагилар такрорланмайди
        self.client.post(url, {"bank": bank.pk, "mode": "random", "count": "20"})
        self.assertEqual(self.logic.questions.filter(source__bank=bank).values("source").distinct().count(), 30)
        # Ортиқча сон рад этилади
        r = self.client.post(url, {"bank": bank.pk, "mode": "random", "count": "5"}, follow=True)
        self.assertContains(r, "гача сон киритинг")
        self.assertEqual(self.logic.questions.filter(source__bank=bank).count(), 30)


class PerCandidateDrawTests(Base):
    def test_each_candidate_gets_own_set(self):
        for i in range(3, 9):  # амалий блокда жами 8 та савол
            q = Question.objects.create(block=self.practical, code=f"МА 0{i}", title=f"V{i}", text=f"<p>{i}</p>",
                                        qtype="text", max_score=10, order=i)
            q.criteria.create(text="x", max_score=2)
        self.practical.questions_per_candidate = 3
        self.practical.save()
        self.assertTrue(self.practical.is_drawn)
        self.assertEqual(self.practical.max_per_candidate, 30)
        self.assertEqual(self.test.max_score_for(Kind.PRACTICAL), 30)

        sets = []
        for n in range(8):
            a = self.register(f"Nomzod{n}", "Test", "", "1990")
            self.client.logout()
            c = self.client_class()
            c.post(reverse("exam_login"), {"login": a.login, "password": a.password})
            attempt = Attempt.objects.get(candidate=a.candidate)
            practical = attempt.answers.filter(question__block=self.practical)
            self.assertEqual(practical.count(), 3)
            self.assertEqual(attempt.answers.filter(question__block=self.logic).count(), 2)  # бошқа блок тўлиқ
            sets.append(frozenset(practical.values_list("question_id", flat=True)))
            r = c.get(reverse("exam_take"))
            self.assertEqual(r.status_code, 200)
        self.assertGreater(len(set(sets)), 1)  # номзодларга ҳар хил тўплам тушган

        # Балл ва якуний натижа номзоднинг ўз саволлари бўйича
        attempt = Attempt.objects.first()
        from . import grading
        self.assertEqual(grading.summary(attempt)["practical"]["max"], 30)

    def test_zero_means_all(self):
        self.assertFalse(self.practical.is_drawn)
        self.assertEqual(self.practical.per_candidate, 2)


class BlockRuleTests(Base):
    def make_bank(self, name, n, kind=Kind.LOGIC):
        bank = QuestionBank.objects.create(name=name, kind=kind)
        for i in range(n):
            q = Question.objects.create(bank=bank, code=f"{name}{i}", title=f"{name} {i}", text=f"<p>{i}</p>", qtype="choice", max_score=1, order=i)
            o = q.options.create(text="ha", order=1); q.options.create(text="yo'q", order=2)
            q.correct_option = o; q.save()
        return bank

    def test_rules_create_draw_and_edit(self):
        b1, b2 = self.make_bank("A", 20), self.make_bank("B", 20)
        self.client.force_login(self.author)
        r = self.client.post(reverse("block_create", args=[self.test.pk]), {
            "title": "Aralash", "kind": "logic", "duration_minutes": 30, "order": 3,
            "rule_bank": [b1.pk, b2.pk], "rule_count": ["5", "3"],
        })
        self.assertEqual(r.status_code, 302)
        block = self.test.blocks.get(title="Aralash")
        self.assertEqual(block.pool_size, 40)
        self.assertEqual(block.per_candidate, 8)
        self.assertEqual(block.max_per_candidate, 8)
        self.assertTrue(block.is_drawn)
        r = self.client.get(reverse("test_detail", args=[self.test.pk]))
        self.assertContains(r, "A: 5 та")

        sets = []
        for n in range(6):
            a = self.register(f"R{n}", "T", "", "1990")
            c = self.client_class()
            c.post(reverse("exam_login"), {"login": a.login, "password": a.password})
            attempt = Attempt.objects.get(candidate=a.candidate)
            got = attempt.answers.filter(question__block=block)
            self.assertEqual(got.filter(question__rule__bank=b1).count(), 5)
            self.assertEqual(got.filter(question__rule__bank=b2).count(), 3)
            sets.append(frozenset(got.values_list("question_id", flat=True)))
        self.assertGreater(len(set(sets)), 1)

    def test_rule_validation_edit_and_remove(self):
        b1, b2 = self.make_bank("A", 4), self.make_bank("B", 4)
        self.client.force_login(self.author)
        r = self.client.post(reverse("block_create", args=[self.test.pk]), {
            "title": "X", "kind": "logic", "duration_minutes": 10, "order": 3, "rule_bank": [b1.pk], "rule_count": ["9"]})
        self.assertContains(r, "1 дан 4 гача")
        self.client.post(reverse("block_create", args=[self.test.pk]), {
            "title": "X", "kind": "logic", "duration_minutes": 10, "order": 3,
            "rule_bank": [b1.pk, b2.pk], "rule_count": ["2", "2"]})
        block = self.test.blocks.get(title="X")
        self.client.post(reverse("block_edit", args=[block.pk]), {
            "title": "X", "kind": "logic", "duration_minutes": 10, "order": 3, "rule_bank": [b1.pk], "rule_count": ["3"]})
        block.refresh_from_db()
        self.assertEqual(block.pool_size, 4)  # B олиб ташланди
        self.assertEqual(block.per_candidate, 3)

    def test_copy_keeps_rules(self):
        b1 = self.make_bank("A", 10)
        self.client.force_login(self.author)
        self.client.post(reverse("block_create", args=[self.test.pk]), {
            "title": "X", "kind": "logic", "duration_minutes": 10, "order": 3, "rule_bank": [b1.pk], "rule_count": ["4"]})
        self.client.post(reverse("test_copy", args=[self.test.pk]))
        copy = Test.objects.exclude(pk=self.test.pk).get()
        nb = copy.blocks.get(title="X")
        self.assertEqual((nb.pool_size, nb.per_candidate), (10, 4))


class BankListTests(Base):
    def test_grouped_by_kind(self):
        for name, kind in [("L2", Kind.LOGIC), ("P1", Kind.PRACTICAL), ("L1", Kind.LOGIC), ("P2", Kind.PRACTICAL)]:
            QuestionBank.objects.create(name=name, kind=kind)
        self.client.force_login(self.author)
        html = self.client.get(reverse("bank_list")).content.decode()
        self.assertEqual(html.count("Мантиқий тест · "), 1)  # битта гуруҳ
        self.assertLess(html.index(">P1<"), html.index(">P2<"))
        self.assertLess(html.index(">P2<"), html.index(">L1<"))
        self.assertLess(html.index(">L1<"), html.index(">L2<"))


class BankMergeTests(BlockRuleTests):
    def test_merge_moves_questions_and_rules(self):
        b1, b2, b3 = self.make_bank("A", 5), self.make_bank("B", 5), self.make_bank("C", 5)
        self.client.force_login(self.author)
        self.client.post(reverse("block_create", args=[self.test.pk]), {
            "title": "X", "kind": "logic", "duration_minutes": 10, "order": 3,
            "rule_bank": [b1.pk, b2.pk], "rule_count": ["2", "3"]})
        block = self.test.blocks.get(title="X")
        self.client.post(reverse("bank_merge", args=[b1.pk]), {"merge": [b2.pk, b3.pk], "name": "Mantiq"})
        b1.refresh_from_db()
        self.assertEqual(b1.name, "Mantiq")
        self.assertEqual(b1.questions.count(), 15)
        self.assertFalse(QuestionBank.objects.filter(pk__in=[b2.pk, b3.pk]).exists())
        rules = list(block.rules.all())
        self.assertEqual([(r.bank_id, r.count) for r in rules], [(b1.pk, 5)])
        self.assertEqual(block.pool_size, 10)
        self.assertEqual(block.per_candidate, 5)
        # Синов: номзод 5 та савол олади
        a = self.register("M", "T", "", "1990")
        self.client_class().post(reverse("exam_login"), {"login": a.login, "password": a.password})
        self.assertEqual(Attempt.objects.get(candidate=a.candidate).answers.filter(question__block=block).count(), 5)


class AssignmentEditTests(Base):
    def test_change_test_and_name(self):
        other = Test.objects.create(title="Boshqa", is_active=True, created_by=self.admin, require_photo=False)
        other.blocks.create(title="B", kind="logic", duration_minutes=5)
        hr = User.objects.create_user("hr", password="x" * 10, role=User.Role.HR)
        self.client.force_login(hr)
        a = self.assignment
        login, pw = a.login, a.password
        r = self.client.post(reverse("assignment_edit", args=[a.pk]), {
            "last_name": "Karimov", "first_name": "Aziz", "middle_name": "Olimovich", "birth": "1996", "test": other.pk})
        self.assertEqual(r.status_code, 302)
        a.refresh_from_db()
        self.assertEqual((a.test_id, a.login, a.password, a.candidate.birth_year), (other.pk, login, pw, 1996))
        self.client.logout()
        self.assertRedirects(self.login_candidate(), reverse("exam_take"))
        self.assertEqual(Attempt.objects.get().test, other)
        # Бошлагандан кейин тест ўзгармайди
        self.client.force_login(hr)
        self.client.post(reverse("assignment_edit", args=[a.pk]), {
            "last_name": "Karimov", "first_name": "Aziz", "middle_name": "Olimovich", "birth": "1996", "test": self.test.pk})
        a.refresh_from_db()
        self.assertEqual(a.test_id, other.pk)
        self.assertContains(self.client.get(reverse("assignment_edit", args=[a.pk])), "тестни ўзгартириб бўлмайди")


class DonePageTests(Base):
    def test_done_returns_to_login(self):
        self.take_test()  # take_test exam_done ни очади
        self.login_candidate()  # якунлангандан кейин
        r = self.client.get(reverse("exam_done"))
        self.assertRedirects(r, reverse("exam_login"))  # сессия тозаланган
        self.assertEqual(self.client.get(reverse("exam_take")).status_code, 302)


class MissingReviewerTests(Base):
    def test_stuck_work_shown_and_fixed_by_single_mode(self):
        self.test.checkers.set([])
        attempt = self.take_test()
        self.assertEqual(attempt.reviews.count(), 0)
        self.client.force_login(self.admin)
        # суперадмин битта экспертни қўлда қўшади
        self.client.post(reverse("result_action", args=[attempt.pk]), {"action": "add_review", "checker": self.c1.pk})
        self.grade(self.c1, attempt, {self.q_ma1: [2] * 5, self.q_ma2: [1] * 5})
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.FINISHED)  # иккинчи эксперт кутилмоқда
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("results")), "1 та эксперт йўқ")
        self.assertContains(self.client.get(reverse("result_detail", args=[attempt.pk])), "Яна 1 та эксперт керак")
        # «Битта эксперт» га ўтказилса — дарҳол баҳоланади
        self.client.force_login(self.author)
        self.client.post(reverse("test_edit", args=[self.test.pk]), {
            "title": self.test.title, "review_mode": "single", "is_active": "on",
            "weight_practical": 55, "weight_logic": 15, "weight_interview": 15, "weight_defense": 15,
            "interview_questions": [self.iv.pk]})
        attempt.refresh_from_db()
        self.assertEqual(attempt.status, Attempt.Status.GRADED)
        self.assertEqual(attempt.answers.get(question=self.q_ma1).score, 10)

    def test_assigning_checkers_fills_missing(self):
        self.test.checkers.set([])
        attempt = self.take_test()
        self.client.force_login(self.admin)
        self.client.post(reverse("test_checkers", args=[self.test.pk]), {"checkers": [self.c1.pk, self.c2.pk]})
        self.assertEqual(attempt.reviews.count(), 2)


class HrResultsTests(Base):
    def test_hr_sees_codes_not_names(self):
        attempt = self.take_test()
        self.grade(attempt.reviews.first().checker, attempt, {self.q_ma1: [2] * 5, self.q_ma2: [1] * 5})
        hr = User.objects.create_user("hr", password="x" * 10, role=User.Role.HR)
        self.client.force_login(hr)
        r = self.client.get(reverse("hr_results"))
        attempt.refresh_from_db()
        self.assertContains(r, attempt.code)
        self.assertContains(r, attempt.get_status_display())
        for secret in ("Karimov", "Aziz", "1995", "photo.jpg"):
            self.assertNotContains(r, secret)
        # Номзодлар рўйхатида код кўринмайди (исм ва кодни боғлаб бўлмасин)
        self.assertNotContains(self.client.get(reverse("candidate_list")), attempt.code)
        # Тўлиқ натижалар ва номзод саҳифаси HR учун ёпиқ
        self.assertEqual(self.client.get(reverse("results")).status_code, 403)
        self.assertEqual(self.client.get(reverse("result_detail", args=[attempt.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("results_export")).status_code, 403)


class CandidateListFilterTests(Base):
    def test_edit_returns_to_same_view_and_filter_is_visible(self):
        self.client.force_login(self.admin)
        a = self.assignment
        r = self.client.get(reverse("assignment_edit", args=[a.pk]) + "?back=q%3DKar")
        self.assertContains(r, 'href="/staff/candidates/?q=Kar"')
        r = self.client.post(reverse("assignment_edit", args=[a.pk]), {
            "back": "", "last_name": "Karimov", "first_name": "Aziz", "middle_name": "Olimovich", "birth": "1995", "test": self.test.pk})
        self.assertRedirects(r, reverse("candidate_list"))
        r = self.client.get(reverse("candidate_list") + f"?test={self.test.pk}")
        self.assertContains(r, "Филтр ёқилган")
        self.assertContains(r, "Барчасини кўрсатиш")
        # Ташқи манзилга йўналтириб бўлмайди
        r = self.client.get(reverse("assignment_edit", args=[a.pk]) + "?back=//evil.com")
        self.assertNotContains(r, "evil.com\"")


class NoInterviewTests(Base):
    def test_final_without_interview_and_defense(self):
        self.client.force_login(self.author)
        data = {"title": self.test.title, "review_mode": "single", "is_active": "on",
                "weight_practical": 80, "weight_logic": 20, "weight_interview": 15, "weight_defense": 0}
        r = self.client.post(reverse("test_edit", args=[self.test.pk]), data)
        self.assertContains(r, "Суҳбат саволлари танланмаган")
        data["weight_interview"] = 0
        data["weight_practical"] = 80
        self.assertEqual(self.client.post(reverse("test_edit", args=[self.test.pk]), data).status_code, 302)
        attempt = self.take_test()
        self.grade(attempt.reviews.first().checker, attempt, {self.q_ma1: [2] * 5, self.q_ma2: [2] * 5})
        attempt.refresh_from_db()
        from . import grading
        s = grading.summary(attempt)
        self.assertIsNone(s["interview"])
        self.assertEqual(s["final"], Decimal("80") + Decimal("10"))  # амалий 20/20 → 80, мантиқий 1/2 → 10


class UserDeleteTests(Base):
    def test_delete_rules(self):
        self.client.force_login(self.admin)
        free = User.objects.create_user("bosh", password="x" * 10, role=User.Role.CHECKER)
        self.client.post(reverse("user_delete", args=[free.pk]))
        self.assertFalse(User.objects.filter(pk=free.pk).exists())
        # Иш текширган эксперт ўчирилмайди
        attempt = self.take_test()
        reviewer = attempt.reviews.first().checker
        self.client.force_login(self.admin)
        r = self.client.post(reverse("user_delete", args=[reviewer.pk]), follow=True)
        self.assertContains(r, "ишни текширган")
        self.assertTrue(User.objects.filter(pk=reviewer.pk).exists())
        # Тест яратган ва ўзини ўчириб бўлмайди
        self.client.post(reverse("user_delete", args=[self.author.pk]))
        self.assertTrue(User.objects.filter(pk=self.author.pk).exists())
        self.client.post(reverse("user_delete", args=[self.admin.pk]))
        self.assertTrue(User.objects.filter(pk=self.admin.pk).exists())
        # Фақат суперадмин
        self.client.force_login(self.author)
        self.assertEqual(self.client.post(reverse("user_delete", args=[self.c3.pk])).status_code, 403)


class ResetAndKeysTests(Base):
    def test_reset_keeps_users(self):
        self.take_test()
        n_users = User.objects.count()
        with self.assertRaises(Exception):
            call_command("reset_tanlov", verbosity=0)
        call_command("reset_tanlov", yes_delete_everything=True, verbosity=0)
        self.assertEqual((Test.objects.count(), Attempt.objects.count(), QuestionBank.objects.count(),
                          Candidate.objects.count()), (0, 0, 0, 0))
        self.assertEqual(User.objects.count(), n_users)

    def test_keys_from_original(self):
        from .importer import parse, use_keys_from
        orig = parse(make_docx(logic_key="Б"))
        edited_doc = make_docx(logic_key="Б")  # калит жадвалисиз нусхани тақлид қиламиз
        edited = parse(edited_doc)
        edited.keys = {}
        ok = use_keys_from(edited, orig)
        self.assertEqual(ok.errors, [])
        # Вариантлар ўзгарган бўлса — хато
        changed = parse(make_docx())
        changed.sections[1].items[0]["options"][0] = ("А", "бошқа")
        self.assertTrue(any("фарқ қилади" in e for e in use_keys_from(changed, orig).errors))
